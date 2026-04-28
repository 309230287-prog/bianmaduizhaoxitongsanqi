from io import BytesIO
from pathlib import Path
from time import monotonic, sleep

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from product_code_mapper.api.app import create_app
from product_code_mapper.model.client import FakeModelClient


def test_api_runs_task_and_exports_excel(tmp_path: Path):
    client = TestClient(create_app(data_dir=tmp_path / "data"))
    client.app.state.task_store._model_client_factory = lambda: FakeModelClient()

    company_response = client.post(
        "/catalog/company/import",
        files={
            "file": (
                "我司商品库.xlsx",
                _workbook_bytes(
                    ["商品编码", "商品名称", "品牌", "规格", "单位"],
                    [["P1", "海天金标生抽", "海天", "500ml", "瓶"]],
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert company_response.status_code == 200
    assert company_response.json()["imported_count"] == 1

    task_response = client.post(
        "/tasks",
        files={
            "file": (
                "客户商品库.xlsx",
                _workbook_bytes(
                    ["商品名称", "品牌", "规格", "单位"],
                    [
                        ["海天金标生抽", "海天", "500ml", "瓶"],
                        ["番茄酱", "", "500g", "瓶"],
                    ],
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert task_response.status_code == 200
    task_id = task_response.json()["task_id"]
    assert task_response.json()["customer_count"] == 2

    _confirm_suggested_fields(client, task_id)
    start_response = client.post(f"/tasks/{task_id}/start")
    assert start_response.status_code == 200
    assert "task_status" in start_response.json()
    assert "run_status" in start_response.json()

    status_payload = _wait_for_task_status(client, task_id, "completed")
    assert status_payload["metrics"]["auto_code_count"] == 1

    export_response = client.get(f"/tasks/{task_id}/export")
    assert export_response.status_code == 200

    workbook = load_workbook(BytesIO(export_response.content))
    assert workbook.sheetnames == ["对照结果总表", "详细证据表", "候选明细表", "统计汇总表"]
    summary_sheet = workbook["对照结果总表"]
    headers = [cell.value for cell in summary_sheet[1]]
    product_code_column = headers.index("我司商品编码") + 1
    assert summary_sheet.cell(row=2, column=product_code_column).value == "P1"


def test_api_can_pause_and_stop_running_task(tmp_path: Path):
    client = TestClient(create_app(data_dir=tmp_path / "data"))
    client.app.state.task_store._model_client_factory = lambda: _SlowFakeModelClient()

    client.post(
        "/catalog/company/import",
        files={
            "file": (
                "我司商品库.xlsx",
                _workbook_bytes(
                    ["商品编码", "商品名称", "品牌", "规格", "单位"],
                    [["P1", "海天金标生抽", "海天", "500ml", "瓶"]],
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    task_response = client.post(
        "/tasks",
        files={
            "file": (
                "客户商品库.xlsx",
                _workbook_bytes(
                    ["商品名称", "品牌", "规格", "单位"],
                    [["海天金标生抽", "海天", "500ml", "瓶"] for _ in range(20)],
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    task_id = task_response.json()["task_id"]

    _confirm_suggested_fields(client, task_id)
    start_response = client.post(f"/tasks/{task_id}/start")
    assert start_response.status_code == 200
    assert start_response.json()["run_status"] == "running"

    pause_response = client.post(f"/tasks/{task_id}/pause")
    assert pause_response.status_code == 200
    paused_payload = _wait_for_run_status(client, task_id, "paused")
    assert paused_payload["can_export"] is False

    stop_response = client.post(f"/tasks/{task_id}/stop")
    assert stop_response.status_code == 200
    stopped_payload = _wait_for_task_status(client, task_id, "stopped")
    assert stopped_payload["can_export"] is False

    export_response = client.get(f"/tasks/{task_id}/export")
    assert export_response.status_code == 409


def test_running_task_status_exposes_live_metrics(tmp_path: Path):
    client = TestClient(create_app(data_dir=tmp_path / "data"))
    client.app.state.task_store._model_client_factory = lambda: _SlowFakeModelClient(delay_seconds=0.08)

    _import_company_catalog(client, [["P1", "海天金标生抽", "海天", "500ml", "瓶"]])
    task_id = _create_customer_task(
        client,
        [["海天金标生抽", "海天", "500ml", "瓶"] for _ in range(20)],
    )
    _confirm_suggested_fields(client, task_id)

    start_response = client.post(f"/tasks/{task_id}/start")
    assert start_response.status_code == 200

    payload = _wait_for_live_metrics(client, task_id)

    assert payload["run_status"] == "running"
    assert payload["metrics"]["total_count"] == 20
    processed = (
        payload["metrics"]["auto_code_count"]
        + payload["metrics"]["manual_review_count"]
        + payload["metrics"]["no_reliable_match_count"]
        + payload["metrics"]["suggested_review_count"]
    )
    assert processed > 0


def test_api_task_start_uses_model_round_planning(tmp_path: Path):
    client = TestClient(create_app(data_dir=tmp_path / "data"))
    model = _PlanningSpyModelClient()
    client.app.state.task_store._model_client_factory = lambda: model

    _import_company_catalog(client, [["P1", "海天金标生抽", "海天", "500ml", "瓶"]])
    task_id = _create_customer_task(client, [["海天金标生抽", "海天", "500ml", "瓶"]])

    _confirm_suggested_fields(client, task_id)
    client.post(f"/tasks/{task_id}/start")
    _wait_for_task_status(client, task_id, "completed")

    assert model.plan_round_calls >= 1


def test_config_status_lists_tasks_after_task_created(tmp_path: Path):
    client = TestClient(create_app(data_dir=tmp_path / "data"))
    task_id = _create_customer_task(client, [["海天金标生抽", "海天", "500ml", "瓶"]])

    response = client.get("/config/status")

    assert response.status_code == 200
    tasks = response.json()["tasks"]
    assert tasks[0]["task_id"] == task_id
    assert tasks[0]["task_status"] == "created"


def test_multiple_excel_tasks_can_reuse_task_local_row_ids(tmp_path: Path):
    client = TestClient(create_app(data_dir=tmp_path / "data"))

    first_task_id = _create_customer_task(client, [["海天金标生抽", "海天", "500ml", "瓶"]])
    second_task_id = _create_customer_task(client, [["李锦记生抽", "李锦记", "500ml", "瓶"]])

    assert first_task_id != second_task_id
    assert client.get(f"/tasks/{first_task_id}/status").json()["customer_count"] == 1
    assert client.get(f"/tasks/{second_task_id}/status").json()["customer_count"] == 1


def test_next_round_does_not_overwrite_human_confirmed_rows(tmp_path: Path):
    client = TestClient(create_app(data_dir=tmp_path / "data"))
    client.app.state.task_store._model_client_factory = lambda: FakeModelClient()

    _import_company_catalog(client, [["P1", "海天金标生抽", "海天", "500ml", "瓶"]])
    task_id = _create_customer_task(client, [["海天金标生抽", "海天", "500ml", "瓶"]])
    _confirm_suggested_fields(client, task_id)
    client.post(f"/tasks/{task_id}/start")
    _wait_for_task_status(client, task_id, "completed")

    export_response = client.get(f"/tasks/{task_id}/export")
    reviewed_excel = _mark_first_row_confirmed(export_response.content)

    _import_company_catalog(client, [
        ["P2", "海天金标生抽", "海天", "500ml", "瓶"],
        ["P1", "海天金标生抽", "海天", "500ml", "瓶"],
    ])
    preview_response = client.post(
        f"/tasks/{task_id}/next-round/preview",
        files={
            "file": (
                "人工审核后.xlsx",
                reviewed_excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview_response.status_code == 200
    assert preview_response.json()["valid"] is True

    next_round_response = client.post(f"/tasks/{task_id}/next-round/start")
    assert next_round_response.status_code == 200

    second_export = client.get(f"/tasks/{task_id}/export")
    workbook = load_workbook(BytesIO(second_export.content))
    sheet = workbook["对照结果总表"]
    headers = [cell.value for cell in sheet[1]]
    code_column = headers.index("我司商品编码") + 1
    assert sheet.cell(row=2, column=code_column).value == "P1"


def _wait_for_task_status(client: TestClient, task_id: str, expected_status: str) -> dict:
    deadline = monotonic() + 5
    last_payload: dict = {}
    while monotonic() < deadline:
        response = client.get(f"/tasks/{task_id}/status")
        assert response.status_code == 200
        last_payload = response.json()
        if last_payload.get("task_status") == expected_status:
            return last_payload
        sleep(0.02)
    raise AssertionError(f"任务未进入 {expected_status} 状态，最后状态: {last_payload}")


def _wait_for_run_status(client: TestClient, task_id: str, expected_status: str) -> dict:
    deadline = monotonic() + 5
    last_payload: dict = {}
    while monotonic() < deadline:
        response = client.get(f"/tasks/{task_id}/status")
        assert response.status_code == 200
        last_payload = response.json()
        if last_payload.get("run_status") == expected_status:
            return last_payload
        sleep(0.02)
    raise AssertionError(f"运行未进入 {expected_status} 状态，最后状态: {last_payload}")


def _wait_for_live_metrics(client: TestClient, task_id: str) -> dict:
    deadline = monotonic() + 3
    last_payload: dict = {}
    while monotonic() < deadline:
        response = client.get(f"/tasks/{task_id}/status")
        assert response.status_code == 200
        last_payload = response.json()
        metrics = last_payload.get("metrics") or {}
        processed = (
            metrics.get("auto_code_count", 0)
            + metrics.get("manual_review_count", 0)
            + metrics.get("no_reliable_match_count", 0)
            + metrics.get("suggested_review_count", 0)
        )
        if last_payload.get("run_status") == "running" and processed > 0:
            return last_payload
        sleep(0.02)
    raise AssertionError(f"运行中没有返回实时指标，最后状态: {last_payload}")


def _confirm_suggested_fields(client: TestClient, task_id: str) -> None:
    suggestions = client.get(f"/tasks/{task_id}/fields/suggestions")
    assert suggestions.status_code == 200
    payload = suggestions.json()
    response = client.post(
        f"/tasks/{task_id}/fields/confirm",
        json={
            "customer_mappings": payload["customer_mappings"],
            "company_mappings": payload["company_mappings"],
        },
    )
    assert response.status_code == 200
    assert response.json()["confirmed"] is True


class _SlowFakeModelClient(FakeModelClient):
    def __init__(self, delay_seconds: float = 0.03):
        self._delay_seconds = delay_seconds

    def compare(self, customer, candidates):
        sleep(self._delay_seconds)
        return super().compare(customer, candidates)


class _PlanningSpyModelClient(FakeModelClient):
    def __init__(self):
        self.plan_round_calls = 0

    def plan_round(self, *args, **kwargs):
        self.plan_round_calls += 1
        return super().plan_round(*args, **kwargs)


def _workbook_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _import_company_catalog(client: TestClient, rows: list[list[str]]) -> None:
    response = client.post(
        "/catalog/company/import",
        files={
            "file": (
                "我司商品库.xlsx",
                _workbook_bytes(["商品编码", "商品名称", "品牌", "规格", "单位"], rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200


def _create_customer_task(client: TestClient, rows: list[list[str]]) -> str:
    response = client.post(
        "/tasks",
        files={
            "file": (
                "客户商品库.xlsx",
                _workbook_bytes(["商品名称", "品牌", "规格", "单位"], rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    return response.json()["task_id"]


def _mark_first_row_confirmed(content: bytes) -> bytes:
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["对照结果总表"]
    headers = [cell.value for cell in sheet[1]]
    review_column = headers.index("人工确认结果") + 1
    sheet.cell(row=2, column=review_column).value = "已确认"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
