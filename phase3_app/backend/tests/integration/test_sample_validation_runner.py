from pathlib import Path

from openpyxl import Workbook, load_workbook

from product_code_mapper.model.client import FakeModelClient
from product_code_mapper.validation.sample_runner import run_sample_validation


def test_sample_validation_runner_exports_limited_real_workbook_shape(tmp_path: Path):
    company_path = tmp_path / "我司商品库.xlsx"
    customer_path = tmp_path / "客户商品库.xlsx"
    output_path = tmp_path / "小批量验证结果.xlsx"
    _write_workbook(
        company_path,
        ["商品编码", "商品名称", "品牌", "规格", "单位"],
        [["P1", "海天金标生抽", "海天", "500ml", "瓶"]],
    )
    _write_workbook(
        customer_path,
        ["商品名称", "品牌", "规格", "单位"],
        [
            ["海天金标生抽", "海天", "500ml", "瓶"],
            ["未知商品", "", "", ""],
        ],
    )

    summary = run_sample_validation(
        company_path=company_path,
        customer_path=customer_path,
        output_path=output_path,
        model_client=FakeModelClient(),
        limit=1,
    )

    assert summary["total_count"] == 1
    assert summary["output_path"] == str(output_path)
    workbook = load_workbook(output_path)
    assert workbook["对照结果总表"].max_row == 2
    assert "候选明细表" in workbook.sheetnames


def test_sample_validation_runner_can_filter_customer_rows_by_keyword(tmp_path: Path):
    company_path = tmp_path / "我司商品库.xlsx"
    customer_path = tmp_path / "客户商品库.xlsx"
    output_path = tmp_path / "按关键词验证结果.xlsx"
    _write_workbook(
        company_path,
        ["商品编码", "商品名称", "品牌", "规格", "单位"],
        [["P1", "海天金标生抽", "海天", "500ml", "瓶"]],
    )
    _write_workbook(
        customer_path,
        ["商品名称", "品牌", "规格", "单位"],
        [
            ["未知商品", "", "", ""],
            ["海天金标生抽", "海天", "500ml", "瓶"],
            ["李锦记生抽", "李锦记", "500ml", "瓶"],
        ],
    )

    summary = run_sample_validation(
        company_path=company_path,
        customer_path=customer_path,
        output_path=output_path,
        model_client=FakeModelClient(),
        limit=10,
        customer_keywords=["海天"],
    )

    assert summary["total_count"] == 1
    workbook = load_workbook(output_path)
    summary_sheet = workbook["对照结果总表"]
    assert summary_sheet.cell(row=2, column=1).value == "海天金标生抽"


def _write_workbook(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
