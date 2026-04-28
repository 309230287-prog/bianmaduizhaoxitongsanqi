from io import BytesIO
from pathlib import Path
from time import monotonic, sleep

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from product_code_mapper.api.app import create_app
from product_code_mapper.model.client import FakeModelClient


def test_field_confirmation_is_required_before_start(tmp_path: Path):
    client = _client_with_fake_model(tmp_path)
    _import_company_catalog(client)
    task_id = _create_customer_task(client)

    suggestions = client.get(f"/tasks/{task_id}/fields/suggestions")
    assert suggestions.status_code == 200
    payload = suggestions.json()
    assert payload["can_start"] is False
    assert payload["customer_mappings"][0]["business_field"] == "商品名称"

    blocked = client.post(f"/tasks/{task_id}/start")
    assert blocked.status_code == 409
    assert "字段确认" in blocked.text

    confirmed = client.post(
        f"/tasks/{task_id}/fields/confirm",
        json={
            "customer_mappings": payload["customer_mappings"],
            "company_mappings": payload["company_mappings"],
        },
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["confirmed"] is True

    started = client.post(f"/tasks/{task_id}/start")
    assert started.status_code == 200
    assert _wait_for_task_status(client, task_id, "completed")["can_export"] is True


def test_missing_real_model_configuration_blocks_task_start(tmp_path: Path):
    client = TestClient(create_app(data_dir=tmp_path / "data"))
    _import_company_catalog(client)
    task_id = _create_customer_task(client)
    _confirm_suggested_fields(client, task_id)

    response = client.post(f"/tasks/{task_id}/start")

    assert response.status_code == 409
    assert "模型未配置" in response.text


def test_manual_item_can_create_independent_task(tmp_path: Path):
    client = _client_with_fake_model(tmp_path)
    _import_company_catalog(client)

    response = client.post(
        "/tasks/manual",
        json={"商品名称": "海天金标生抽", "品牌": "海天", "规格": "500ml", "单位": "瓶"},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["customer_count"] == 1
    task_id = payload["task_id"]

    _confirm_suggested_fields(client, task_id)
    client.post(f"/tasks/{task_id}/start")
    status = _wait_for_task_status(client, task_id, "completed")
    assert status["metrics"]["auto_code_count"] == 1


def test_multiple_manual_tasks_can_reuse_task_local_manual_row_id(tmp_path: Path):
    client = _client_with_fake_model(tmp_path)

    first = client.post("/tasks/manual", json={"商品名称": "海天金标生抽"})
    second = client.post("/tasks/manual", json={"商品名称": "李锦记生抽"})

    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["task_id"] != second.json()["task_id"]


def test_export_contains_candidate_detail_sheet_and_internal_link(tmp_path: Path):
    client = _client_with_fake_model(tmp_path)
    _import_company_catalog(client)
    task_id = _create_customer_task(client)
    _confirm_suggested_fields(client, task_id)
    client.post(f"/tasks/{task_id}/start")
    _wait_for_task_status(client, task_id, "completed")

    response = client.get(f"/tasks/{task_id}/export")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert "候选明细表" in workbook.sheetnames
    summary = workbook["对照结果总表"]
    headers = [cell.value for cell in summary[1]]
    candidate_col = headers.index("备选候选") + 1
    candidate_cell = summary.cell(row=2, column=candidate_col)
    assert candidate_cell.value
    assert candidate_cell.hyperlink is not None
    assert "候选明细表" in candidate_cell.hyperlink.target

    candidate_sheet = workbook["候选明细表"]
    assert candidate_sheet.max_row >= 2
    candidate_row_ids = [candidate_sheet.cell(row=row, column=1).value for row in range(2, candidate_sheet.max_row + 1)]
    assert "row-2" in candidate_row_ids


def test_invalid_customer_excel_returns_chinese_400_with_cors(tmp_path: Path):
    client = TestClient(create_app(data_dir=tmp_path / "data"), raise_server_exceptions=False)

    response = client.post(
        "/tasks",
        files={"file": ("坏文件.txt", b"not an excel", "text/plain")},
        headers={"Origin": "http://tauri.localhost"},
    )

    assert response.status_code == 400
    assert response.headers["access-control-allow-origin"] == "http://tauri.localhost"
    assert "Excel" in response.json()["detail"]


def test_completed_task_can_export_after_app_restart(tmp_path: Path):
    data_dir = tmp_path / "data"
    client = _client_with_fake_model(tmp_path, data_dir=data_dir)
    _import_company_catalog(client)
    task_id = _create_customer_task(client)
    _confirm_suggested_fields(client, task_id)
    client.post(f"/tasks/{task_id}/start")
    _wait_for_task_status(client, task_id, "completed")

    restarted = _client_with_fake_model(tmp_path, data_dir=data_dir)
    status = restarted.get(f"/tasks/{task_id}/status")
    export_response = restarted.get(f"/tasks/{task_id}/export")

    assert status.status_code == 200
    assert status.json()["can_export"] is True
    assert export_response.status_code == 200
    workbook = load_workbook(BytesIO(export_response.content))
    assert "对照结果总表" in workbook.sheetnames


def test_completed_task_is_not_restarted_when_start_called_again(tmp_path: Path):
    client = _client_with_fake_model(tmp_path)
    _import_company_catalog(client)
    task_id = _create_customer_task(client)
    _confirm_suggested_fields(client, task_id)
    client.post(f"/tasks/{task_id}/start")
    _wait_for_task_status(client, task_id, "completed")

    second_start = client.post(f"/tasks/{task_id}/start")

    assert second_start.status_code == 200
    assert second_start.json()["task_status"] == "completed"
    assert second_start.json()["run_status"] == "completed"
    assert second_start.json()["can_export"] is True


def test_config_status_lists_completed_tasks_after_app_restart(tmp_path: Path):
    data_dir = tmp_path / "data"
    client = _client_with_fake_model(tmp_path, data_dir=data_dir)
    _import_company_catalog(client)
    task_id = _create_customer_task(client)
    _confirm_suggested_fields(client, task_id)
    client.post(f"/tasks/{task_id}/start")
    _wait_for_task_status(client, task_id, "completed")

    restarted = _client_with_fake_model(tmp_path, data_dir=data_dir)
    status = restarted.get("/config/status")

    assert status.status_code == 200
    tasks = status.json()["tasks"]
    assert any(task["task_id"] == task_id and task["can_export"] for task in tasks)


def _client_with_fake_model(tmp_path: Path, data_dir: Path | None = None) -> TestClient:
    client = TestClient(create_app(data_dir=data_dir or tmp_path / "data"))
    client.app.state.task_store._model_client_factory = lambda: FakeModelClient()
    return client


def _import_company_catalog(client: TestClient) -> None:
    response = client.post(
        "/catalog/company/import",
        files={
            "file": (
                "我司商品库.xlsx",
                _workbook_bytes(
                    ["商品编码", "商品名称", "品牌", "规格", "单位"],
                    [
                        ["P1", "海天金标生抽", "海天", "500ml", "瓶"],
                        ["P2", "海天生抽", "海天", "1.9L", "桶"],
                    ],
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200


def _create_customer_task(client: TestClient) -> str:
    response = client.post(
        "/tasks",
        files={
            "file": (
                "客户商品库.xlsx",
                _workbook_bytes(
                    ["商品名称", "品牌", "规格", "单位"],
                    [
                        ["海天金标生抽", "海天", "500ml", "瓶"],
                        ["海天生抽", "海天", "1.9L", "桶"],
                    ],
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    return response.json()["task_id"]


def _confirm_suggested_fields(client: TestClient, task_id: str) -> None:
    suggestions = client.get(f"/tasks/{task_id}/fields/suggestions")
    assert suggestions.status_code == 200
    payload = suggestions.json()
    response = client.post(
        f"/tasks/{task_id}/fields/confirm",
        json={
            "customer_mappings": payload["customer_mappings"],
            "company_mappings": payload["company_mappings"],
        },
    )
    assert response.status_code == 200


def _wait_for_task_status(client: TestClient, task_id: str, expected_status: str) -> dict:
    deadline = monotonic() + 5
    last_payload: dict = {}
    while monotonic() < deadline:
        response = client.get(f"/tasks/{task_id}/status")
        assert response.status_code == 200
        last_payload = response.json()
        if last_payload.get("task_status") == expected_status:
            return last_payload
        sleep(0.02)
    raise AssertionError(f"任务未进入 {expected_status} 状态，最后状态: {last_payload}")


def _workbook_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
