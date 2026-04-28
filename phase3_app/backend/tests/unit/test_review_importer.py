from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from product_code_mapper.excel.review_importer import import_reviewed_excel


def _make_review_workbook(rows: list[list[str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "对照结果总表"
    ws.append(["系统任务行ID", "原始行号", "对照状态", "我司商品编码",
                "我司商品名称", "是否需要人工审核", "人工确认结果", "人工审核备注"])
    for row in rows:
        ws.append(row)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_import_confirmed_rows():
    data = _make_review_workbook([
        ["row-1", "2", "自动落码", "P1", "生抽", "否", "已确认", ""],
        ["row-2", "3", "必须人工审核", "", "", "是", "已确认", "人工核对通过"],
    ])
    company_codes = frozenset(["P1", "P2"])

    result = import_reviewed_excel(BytesIO(data), company_codes)

    assert result.is_valid
    assert result.confirmed_count == 2


def test_import_missing_sheet():
    wb = Workbook()
    wb.active.title = "其他表"
    stream = BytesIO()
    wb.save(stream)

    result = import_reviewed_excel(BytesIO(stream.getvalue()), frozenset())
    assert not result.is_valid
    assert "缺少'对照结果总表'" in result.errors[0]


def test_import_missing_task_row_id_column():
    wb = Workbook()
    ws = wb.active
    ws.title = "对照结果总表"
    ws.append(["其他列", "另一列"])
    stream = BytesIO()
    wb.save(stream)

    result = import_reviewed_excel(BytesIO(stream.getvalue()), frozenset())
    assert not result.is_valid
    assert "系统任务行ID" in result.errors[0]


def test_import_detects_invalid_company_code():
    data = _make_review_workbook([
        ["row-1", "2", "", "P99", "", "", "已确认修改", "改选了P99"],
    ])
    company_codes = frozenset(["P1"])

    result = import_reviewed_excel(BytesIO(data), company_codes)
    assert len(result.errors) >= 1
    assert "P99" in result.errors[0]


def test_import_detects_invalid_confirmed_company_code():
    data = _make_review_workbook([
        ["row-1", "2", "", "P99", "", "", "已确认", "人工确认了不存在的编码"],
    ])
    company_codes = frozenset(["P1"])

    result = import_reviewed_excel(BytesIO(data), company_codes)

    assert not result.is_valid
    assert "P99" in result.errors[0]


def test_import_with_no_match_rows():
    data = _make_review_workbook([
        ["row-1", "2", "", "", "", "", "无匹配", "确实找不到"],
    ])
    company_codes = frozenset(["P1"])

    result = import_reviewed_excel(BytesIO(data), company_codes)
    assert result.is_valid
    assert result.no_match_count == 1
