from pathlib import Path

from openpyxl import load_workbook

from product_code_mapper.domain.models import Candidate, CompanyProduct, CustomerItem, MatchResult
from product_code_mapper.domain.statuses import MatchStatus
from product_code_mapper.excel.exporter import export_run_result
from product_code_mapper.runs.engine import MatchRunResult
from product_code_mapper.runs.metrics import RunMetrics


def test_export_keeps_customer_columns_and_appends_result_columns(tmp_path: Path):
    output_path = tmp_path / "对照结果.xlsx"
    customer = CustomerItem(
        row_id="row-2",
        original_row_index=2,
        fields={"商品名称": "海天金标生抽", "规格": "500ml", "单位": "瓶"},
    )
    product = CompanyProduct(
        code="P1",
        name="海天金标生抽",
        brand="海天",
        spec="500ml",
        unit="瓶",
    )
    result = MatchRunResult(
        metrics=RunMetrics(
            total_count=1,
            auto_code_count=1,
            returned_to_nature_count=0,
            hard_case_count=0,
        ),
        audit_entries=[],
        row_results=[
            MatchResult(
                row_id="row-2",
                status=MatchStatus.AUTO_CODE,
                selected_candidate=Candidate(product=product, score=100),
                reason_summary="品牌、品名、规格、单位一致",
                evidence_summary="系统安全门已通过",
                risk_summary="无",
            )
        ],
        customer_items=[customer],
    )

    export_run_result(result, output_path)

    workbook = load_workbook(output_path)
    assert workbook.sheetnames == ["对照结果总表", "详细证据表", "统计汇总表"]

    summary_sheet = workbook["对照结果总表"]
    headers = [cell.value for cell in summary_sheet[1]]
    assert headers[:3] == ["商品名称", "规格", "单位"]
    assert "系统任务行ID" in headers
    assert headers[3:9] == ["系统任务行ID", "原始行号", "对照状态", "我司商品编码", "我司商品名称", "我司品牌"]
    assert summary_sheet["A2"].value == "海天金标生抽"
    assert summary_sheet["D2"].value == "row-2"
    assert summary_sheet["F2"].value == "自动落码"
    assert summary_sheet["G2"].value == "P1"
