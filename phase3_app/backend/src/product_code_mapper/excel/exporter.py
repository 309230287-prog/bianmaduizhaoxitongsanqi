from __future__ import annotations

from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import Workbook

from product_code_mapper.domain.models import CustomerItem, MatchResult
from product_code_mapper.domain.statuses import MatchStatus
from product_code_mapper.runs.engine import MatchRunResult


APPENDED_RESULT_HEADERS = [
    "系统任务行ID",
    "原始行号",
    "对照状态",
    "我司商品编码",
    "我司商品名称",
    "我司品牌",
    "我司规格",
    "我司单位",
    "推荐理由",
    "证据对齐摘要",
    "风险提示",
    "备选候选",
    "是否需要人工审核",
    "人工确认结果",
    "人工审核备注",
]


def export_run_result(result: MatchRunResult, output_path: str | Path | BinaryIO) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "对照结果总表"
    evidence_sheet = workbook.create_sheet("详细证据表")
    candidate_sheet = workbook.create_sheet("候选明细表")
    metrics_sheet = workbook.create_sheet("统计汇总表")

    candidate_links = _write_candidate_sheet(candidate_sheet, result)
    _write_summary_sheet(summary_sheet, result, candidate_links)
    _write_evidence_sheet(evidence_sheet, result)
    _write_metrics_sheet(metrics_sheet, result)

    if isinstance(output_path, str | Path):
        workbook.save(Path(output_path))
        return
    workbook.save(output_path)


def _write_summary_sheet(sheet, result: MatchRunResult, candidate_links: dict[str, str]) -> None:
    customer_headers = _customer_headers(result.customer_items)
    sheet.append(customer_headers + APPENDED_RESULT_HEADERS)
    candidate_col = len(customer_headers) + APPENDED_RESULT_HEADERS.index("备选候选") + 1

    if isinstance(result.row_results, dict):
        results_by_row_id = result.row_results
    else:
        results_by_row_id = {row_result.row_id: row_result for row_result in result.row_results}
    for item in result.customer_items:
        row_result = results_by_row_id.get(item.row_id)
        sheet.append(
            [item.fields.get(header, "") for header in customer_headers]
            + _result_values(item, row_result)
        )
        link = candidate_links.get(item.row_id)
        if link:
            cell = sheet.cell(row=sheet.max_row, column=candidate_col)
            cell.value = "查看候选明细"
            cell.hyperlink = link
            cell.style = "Hyperlink"


def _write_evidence_sheet(sheet, result: MatchRunResult) -> None:
    sheet.append(["任务行ID", "对照状态", "推荐理由", "证据对齐摘要", "风险提示"])
    row_results_list = (
        list(result.row_results.values())
        if isinstance(result.row_results, dict)
        else result.row_results
    )
    for row_result in row_results_list:
        sheet.append(
            [
                row_result.row_id,
                row_result.status.value,
                row_result.reason_summary,
                row_result.evidence_summary,
                row_result.risk_summary,
            ]
        )


def _write_candidate_sheet(sheet, result: MatchRunResult) -> dict[str, str]:
    sheet.append(["任务行ID", "排名", "我司商品编码", "我司商品名称", "我司品牌", "我司规格", "我司单位", "候选分", "命中词"])
    links: dict[str, str] = {}
    row_results_list = (
        list(result.row_results.values())
        if isinstance(result.row_results, dict)
        else result.row_results
    )
    for row_result in row_results_list:
        candidates = _candidate_summaries_for_export(row_result)
        if not candidates:
            continue
        first_row = sheet.max_row + 1
        links[row_result.row_id] = f"#'候选明细表'!A{first_row}"
        for candidate in candidates:
            sheet.append([
                row_result.row_id,
                candidate.get("rank", ""),
                candidate.get("code", ""),
                candidate.get("name", ""),
                candidate.get("brand", ""),
                candidate.get("spec", ""),
                candidate.get("unit", ""),
                candidate.get("score", ""),
                "、".join(str(term) for term in candidate.get("matched_terms", [])),
            ])
    return links


def _write_metrics_sheet(sheet, result: MatchRunResult) -> None:
    metrics = result.metrics
    sheet.append(["指标", "数量"])
    sheet.append(["总客户行数", metrics.total_count])
    sheet.append(["自动落码", metrics.auto_code_count])
    sheet.append(["自动落码但有表达差异", metrics.auto_code_with_diff_count])
    sheet.append(["建议落码待确认", metrics.suggested_review_count])
    sheet.append(["必须人工审核", metrics.manual_review_count])
    sheet.append(["未找到可靠匹配", metrics.no_reliable_match_count])
    sheet.append(["回到大自然池", metrics.returned_to_nature_count])
    sheet.append(["疑难池", metrics.hard_case_count])
    sheet.append(["总轮次", metrics.total_rounds])


def _customer_headers(customer_items: list[CustomerItem]) -> list[str]:
    headers: list[str] = []
    for item in customer_items:
        for header in item.fields:
            if header not in headers:
                headers.append(header)
    return headers


def _result_values(item: CustomerItem, row_result: MatchResult | None) -> list[Any]:
    if row_result is None:
        return [item.row_id, item.original_row_index] + [""] * (len(APPENDED_RESULT_HEADERS) - 2)

    product = row_result.selected_candidate.product if row_result.selected_candidate else None
    candidate_summary = _candidate_summary_text(row_result)
    return [
        item.row_id,
        item.original_row_index,
        row_result.status.value,
        product.code if product else "",
        product.name if product else "",
        product.brand if product else "",
        product.spec if product else "",
        product.unit if product else "",
        row_result.reason_summary,
        row_result.evidence_summary,
        row_result.risk_summary,
        candidate_summary,
        _needs_review(row_result.status),
        "",
        "",
    ]


def _needs_review(status: MatchStatus) -> str:
    if status in {MatchStatus.AUTO_CODE, MatchStatus.AUTO_CODE_WITH_DIFFERENCE}:
        return "否"
    return "是"


def _candidate_summaries_for_export(row_result: MatchResult) -> list[dict[str, Any]]:
    candidates = row_result.audit.get("candidate_summaries", [])
    if isinstance(candidates, list) and candidates:
        return [candidate for candidate in candidates if isinstance(candidate, dict)]
    if row_result.selected_candidate is None:
        return []
    product = row_result.selected_candidate.product
    return [{
        "rank": 1,
        "code": product.code,
        "name": product.name,
        "brand": product.brand,
        "spec": product.spec,
        "unit": product.unit,
        "score": row_result.selected_candidate.score,
        "matched_terms": row_result.selected_candidate.matched_terms,
    }]


def _candidate_summary_text(row_result: MatchResult) -> str:
    candidates = _candidate_summaries_for_export(row_result)
    if not candidates:
        return ""
    return "；".join(
        f"{candidate.get('rank', '')}. {candidate.get('code', '')} {candidate.get('name', '')} "
        f"{candidate.get('spec', '')} {candidate.get('unit', '')}".strip()
        for candidate in candidates[:5]
    )
