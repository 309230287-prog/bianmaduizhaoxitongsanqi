from __future__ import annotations

from pathlib import Path
from typing import Any, BinaryIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from product_code_mapper.domain.models import CompanyProduct, CustomerItem


WorkbookSource = str | Path | BinaryIO


class ExcelImportError(ValueError):
    pass


def import_customer_items(workbook_path: WorkbookSource) -> list[CustomerItem]:
    workbook = _load_workbook(workbook_path)
    worksheet = workbook.active
    if worksheet is None:
        raise ExcelImportError("Excel 文件不包含任何工作表")
    rows = worksheet.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        raise ExcelImportError("Excel 文件没有表头行，无法读取")

    headers = [_clean_cell(value) for value in header_row]
    if not any(headers):
        raise ExcelImportError("无法识别 Excel 表头，可能是空文件或格式不正确")
    items: list[CustomerItem] = []

    for row_index, row in enumerate(rows, start=2):
        fields = {
            header: _clean_cell(value)
            for header, value in zip(headers, row, strict=False)
            if header
        }
        if _is_empty_row(fields):
            continue
        items.append(
            CustomerItem(
                row_id=f"row-{row_index}",
                original_row_index=row_index,
                fields=fields,
            )
        )

    workbook.close()
    return items


def import_company_products(workbook_path: WorkbookSource) -> list[CompanyProduct]:
    workbook = _load_workbook(workbook_path)
    worksheet = workbook.active
    if worksheet is None:
        raise ExcelImportError("Excel 文件不包含任何工作表")
    rows = worksheet.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        raise ExcelImportError("Excel 文件没有表头行，无法读取")

    headers = [_clean_cell(value) for value in header_row]
    products: list[CompanyProduct] = []

    for row in rows:
        fields = {
            header: _clean_cell(value)
            for header, value in zip(headers, row, strict=False)
            if header
        }
        if _is_empty_row(fields):
            continue
        products.append(
            CompanyProduct(
                code=_first_present(fields, ["商品编码", "编码", "SPUID", "spu_id"]),
                name=_first_present(fields, ["SPU名称（可修改）", "SPU名称", "商品名称", "商品名", "品名", "名称"]),
                brand=_first_present(fields, ["品牌", "品牌名称"]),
                spec=_first_present(fields, ["SPU描述（可修改）", "SPU描述", "规格", "规格型号"]),
                unit=_first_present(fields, ["SPU基本单位", "单位", "计量单位"]),
                package=_first_present(fields, ["包装", "包装单位"]),
                extra_fields=fields,
            )
        )

    workbook.close()
    return products


def _clean_cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _is_empty_row(fields: dict[str, str]) -> bool:
    return not any(value for value in fields.values())


def _normalize_source(workbook_path: WorkbookSource) -> Path | BinaryIO:
    if isinstance(workbook_path, str | Path):
        return Path(workbook_path)
    return workbook_path


def _load_workbook(workbook_path: WorkbookSource):
    try:
        return load_workbook(_normalize_source(workbook_path), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ExcelImportError("无法读取 Excel 文件，请确认上传的是 .xlsx 文件") from exc


def _first_present(fields: dict[str, str], names: list[str]) -> str:
    for name in names:
        value = fields.get(name, "")
        if value:
            return value
    return ""
