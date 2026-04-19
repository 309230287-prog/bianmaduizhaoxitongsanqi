import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from product_matcher_phase2.trial_diagnostics import (  # noqa: E402
    diagnose_trial_results_xlsx,
    render_trial_diagnostics_markdown,
)


class TrialDiagnosticsTests(unittest.TestCase):
    def _write_row(self, worksheet, values):
        worksheet.append(
            [
                values.get("sample_id", ""),
                values.get("sample_group", ""),
                values.get("expected_company_code", ""),
                values.get("expected_result_status", ""),
                values.get("parsed_result_status", ""),
                values.get("status_matches_expected", False),
                values.get("parse_ok", False),
                values.get("selected_candidate_id", ""),
                values.get("selected_company_code", ""),
                values.get("can_auto_code", False),
                values.get("evidence_summary", ""),
                values.get("manual_review_reason", ""),
                values.get("error_message", ""),
                values.get("raw_model_output", ""),
            ]
        )

    def test_diagnose_trial_results_xlsx_classifies_root_causes(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "model_trial_results"
        worksheet.append(
            [
                "sample_id",
                "sample_group",
                "expected_company_code",
                "expected_result_status",
                "parsed_result_status",
                "status_matches_expected",
                "parse_ok",
                "selected_candidate_id",
                "selected_company_code",
                "can_auto_code",
                "evidence_summary",
                "manual_review_reason",
                "error_message",
                "raw_model_output",
            ]
        )
        self._write_row(
            worksheet,
            {
                "sample_id": "S1",
                "raw_model_output": "{not-json",
            },
        )
        self._write_row(
            worksheet,
            {
                "sample_id": "S2",
                "raw_model_output": json.dumps(
                    {
                        "customer_semantic_summary": "x",
                        "candidate_assessments": [
                            {
                                "candidate_id": "K1",
                                "same_business_identity": True,
                                "risk_flags": ["bad_flag"],
                                "summary": "",
                            }
                        ],
                        "selected_candidate_id": "K1",
                        "result_status": "strong_auto_code",
                        "risk_flags": [],
                        "evidence_summary": "",
                        "manual_review_reason": "",
                        "can_auto_code": True,
                    },
                    ensure_ascii=False,
                ),
            },
        )
        self._write_row(
            worksheet,
            {
                "sample_id": "S3",
                "raw_model_output": json.dumps(
                    {
                        "customer_semantic_summary": "x",
                        "candidate_assessments": [
                            {
                                "candidate_id": "K1",
                                "same_business_identity": True,
                                "risk_flags": [],
                                "summary": "",
                            }
                        ],
                        "selected_candidate_id": None,
                        "result_status": "strong_auto_code",
                        "risk_flags": [],
                        "evidence_summary": "",
                        "manual_review_reason": "",
                        "can_auto_code": True,
                    },
                    ensure_ascii=False,
                ),
            },
        )
        self._write_row(
            worksheet,
            {
                "sample_id": "S4",
                "raw_model_output": "",
                "error_message": "模型调用失败: timeout",
            },
        )
        self._write_row(
            worksheet,
            {
                "sample_id": "S5",
                "raw_model_output": "   ",
            },
        )

        with tempfile.TemporaryDirectory() as directory_name:
            path = Path(directory_name) / "trial_results.xlsx"
            workbook.save(path)
            workbook.close()

            summary = diagnose_trial_results_xlsx(path)

        self.assertEqual(summary["total_count"], 5)
        self.assertEqual(summary["acceptance_summary"]["status_match_count"], 0)
        self.assertEqual(summary["acceptance_summary"]["selected_code_match_count"], 0)
        self.assertEqual(summary["acceptance_summary"]["selected_code_mismatch_count"], 0)
        self.assertEqual(summary["acceptance_summary"]["unsafe_auto_code_count"], 0)
        self.assertEqual(
            summary["category_counts"],
            {
                "invalid_json": 1,
                "schema_validation_error": 1,
                "business_rule_validation_error": 1,
                "model_call_error": 1,
                "empty_output": 1,
                "schema_valid_after_normalization": 0,
                "unknown": 0,
            },
        )
        self.assertEqual(summary["representative_samples"]["invalid_json"][0]["sample_id"], "S1")
        self.assertEqual(
            summary["representative_samples"]["business_rule_validation_error"][0]["sample_id"],
            "S3",
        )

    def test_diagnose_trial_results_recomputes_current_parser_metrics_from_raw_output(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "model_trial_results"
        worksheet.append(
            [
                "sample_id",
                "sample_group",
                "expected_company_code",
                "expected_result_status",
                "parsed_result_status",
                "status_matches_expected",
                "parse_ok",
                "selected_candidate_id",
                "selected_company_code",
                "can_auto_code",
                "evidence_summary",
                "manual_review_reason",
                "error_message",
                "raw_model_output",
            ]
        )
        self._write_row(
            worksheet,
            {
                "sample_id": "S6",
                "expected_company_code": "C000001",
                "expected_result_status": "strong_auto_code",
                "parsed_result_status": "model_error",
                "status_matches_expected": False,
                "parse_ok": False,
                "raw_model_output": json.dumps(
                    {
                        "customer_semantic_summary": "客户商品与候选一致。",
                        "candidate_assessments": [
                            {
                                "candidate_id": "K000001",
                                "same_business_identity": True,
                                "risk_flags": [],
                                "summary": "候选完整覆盖客户商品。",
                            }
                        ],
                        "selected_candidate_id": "K000001",
                        "result_status": "strong_auto_code",
                        "risk_flags": [],
                        "evidence_summary": "当前 parser 可解析并命中。",
                        "manual_review_reason": "",
                        "can_auto_code": True,
                    },
                    ensure_ascii=False,
                ),
            },
        )

        with tempfile.TemporaryDirectory() as directory_name:
            directory = Path(directory_name)
            results_path = directory / "trial_results.xlsx"
            input_path = directory / "trial_inputs.jsonl"
            input_path.write_text(
                json.dumps(
                    {
                        "sample_id": "S6",
                        "sample_group": "strong_auto",
                        "expected_company_code": "C000001",
                        "expected_result_status": "strong_auto_code",
                        "payload": {
                            "candidate_products": [
                                {
                                    "candidate_id": "K000001",
                                    "product": {
                                        "code": "C000001",
                                        "name": "测试商品",
                                    },
                                }
                            ]
                        },
                    },
                    ensure_ascii=False,
                )
                + "\n",
                encoding="utf-8",
            )
            workbook.save(results_path)
            workbook.close()

            summary = diagnose_trial_results_xlsx(results_path, trial_input_path=input_path)

        self.assertEqual(summary["acceptance_summary"]["status_match_count"], 0)
        self.assertEqual(summary["current_parser_summary"]["current_parser_valid_count"], 1)
        self.assertEqual(summary["current_parser_summary"]["current_status_match_count"], 1)
        self.assertEqual(summary["current_parser_summary"]["current_selected_code_match_count"], 1)
        self.assertEqual(summary["current_parser_summary"]["current_selected_code_mismatch_count"], 0)
        self.assertEqual(summary["current_parser_summary"]["current_unsafe_auto_code_count"], 0)

    def test_current_parser_summary_applies_payload_safety_gate(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(
            [
                "sample_id",
                "sample_group",
                "expected_company_code",
                "expected_result_status",
                "parsed_result_status",
                "status_matches_expected",
                "parse_ok",
                "selected_candidate_id",
                "selected_company_code",
                "can_auto_code",
                "evidence_summary",
                "manual_review_reason",
                "error_message",
                "raw_model_output",
            ]
        )
        self._write_row(
            worksheet,
            {
                "sample_id": "S7",
                "sample_group": "manual_or_unmatched",
                "expected_company_code": "C000001",
                "expected_result_status": "suggested_code",
                "parsed_result_status": "strong_auto_code",
                "status_matches_expected": False,
                "parse_ok": True,
                "selected_candidate_id": "K000001",
                "selected_company_code": "C000001",
                "can_auto_code": True,
                "raw_model_output": json.dumps(
                    {
                        "customer_semantic_summary": "客户有明确规格。",
                        "candidate_assessments": [
                            {
                                "candidate_id": "K000001",
                                "same_business_identity": True,
                                "risk_flags": [],
                                "summary": "模型误判可自动。",
                            }
                        ],
                        "selected_candidate_id": "K000001",
                        "result_status": "strong_auto_code",
                        "risk_flags": [],
                        "evidence_summary": "模型忽略规格缺失。",
                        "manual_review_reason": "",
                        "can_auto_code": True,
                    },
                    ensure_ascii=False,
                ),
            },
        )

        with tempfile.TemporaryDirectory() as directory_name:
            directory = Path(directory_name)
            results_path = directory / "trial_results.xlsx"
            input_path = directory / "trial_inputs.jsonl"
            input_path.write_text(
                json.dumps(
                    {
                        "sample_id": "S7",
                        "sample_group": "manual_or_unmatched",
                        "expected_company_code": "C000001",
                        "expected_result_status": "suggested_code",
                        "payload": {
                            "customer_record": {"mapped_fields": {"spec": "[1*500g]"}},
                            "candidate_products": [
                                {
                                    "candidate_id": "K000001",
                                    "candidate_evidence": {"spec_tokens": []},
                                    "product": {"code": "C000001", "name": "测试商品"},
                                }
                            ],
                        },
                    },
                    ensure_ascii=False,
                )
                + "\n",
                encoding="utf-8",
            )
            workbook.save(results_path)
            workbook.close()

            summary = diagnose_trial_results_xlsx(results_path, trial_input_path=input_path)

        self.assertEqual(summary["current_parser_summary"]["current_status_match_count"], 1)
        self.assertEqual(summary["current_parser_summary"]["current_unsafe_auto_code_count"], 0)

    def test_render_trial_diagnostics_markdown_includes_summary_and_samples(self) -> None:
        summary = {
            "source_path": "samples/phase2/model_trial_results_deepseek_v0.1.xlsx",
            "total_count": 2,
            "category_counts": {
                "invalid_json": 1,
                "schema_validation_error": 1,
                "business_rule_validation_error": 0,
                "model_call_error": 0,
                "empty_output": 0,
                "schema_valid_after_normalization": 0,
                "unknown": 0,
            },
            "representative_samples": {
                "invalid_json": [{"sample_id": "S1", "reason": "JSON 解析失败"}],
                "schema_validation_error": [{"sample_id": "S2", "reason": "枚举值不合法"}],
                "business_rule_validation_error": [],
                "model_call_error": [],
                "empty_output": [],
                "schema_valid_after_normalization": [],
                "unknown": [],
            },
            "acceptance_summary": {
                "status_match_count": 1,
                "selected_code_match_count": 1,
                "selected_code_mismatch_count": 0,
                "auto_code_count": 1,
                "unsafe_auto_code_count": 0,
            },
            "current_parser_summary": {
                "current_parser_valid_count": 1,
                "current_status_match_count": 1,
                "current_selected_code_match_count": 1,
                "current_selected_code_mismatch_count": 0,
                "current_unsafe_auto_code_count": 0,
            },
        }

        markdown = render_trial_diagnostics_markdown(summary)

        self.assertIn("model_trial_results_deepseek_v0.1.xlsx", markdown)
        self.assertIn("invalid_json", markdown)
        self.assertIn("S1", markdown)
        self.assertIn("Acceptance Summary", markdown)
        self.assertIn("unsafe_auto_code_count", markdown)
        self.assertIn("Current Parser Summary", markdown)
        self.assertIn("current_parser_valid_count", markdown)

    def test_diagnose_real_deepseek_trial_results_sample(self) -> None:
        sample_path = ROOT / "samples" / "phase2" / "model_trial_results_deepseek_v0.1.xlsx"

        summary = diagnose_trial_results_xlsx(sample_path)

        self.assertEqual(summary["total_count"], 10)
        self.assertEqual(summary["category_counts"]["schema_validation_error"], 0)
        self.assertEqual(summary["category_counts"]["business_rule_validation_error"], 1)
        self.assertEqual(summary["category_counts"]["schema_valid_after_normalization"], 8)
        self.assertEqual(summary["category_counts"]["model_call_error"], 1)
        self.assertEqual(summary["representative_samples"]["business_rule_validation_error"][0]["sample_id"], "GS0001")
        self.assertEqual(summary["representative_samples"]["model_call_error"][0]["sample_id"], "GS0021")


if __name__ == "__main__":
    unittest.main()
