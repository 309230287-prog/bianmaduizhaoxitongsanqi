import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from product_matcher.models import MatchCandidate, MatchResult
from product_matcher.services.exporter import build_export_workbook, suggested_export_name


class ExporterTests(unittest.TestCase):
    def test_export_appends_result_columns(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["商品名称", "编号"])
        sheet.append(["海天酱油", "001"])

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            temp_path = Path(tmp.name)
        try:
            workbook.save(temp_path)
            result = MatchResult(
                customer_row_no=2,
                customer_name='海天酱油',
                customer_brand='海天',
                customer_spec='500ml',
                customer_unit='瓶',
                customer_category='调味品',
                match_status='suggested',
                top_score=77.5,
                summary='名称接近；品牌一致',
                candidates=[
                    MatchCandidate(
                        company_row_no=9,
                        company_code='C001',
                        company_name='海天酱油500ml',
                        company_brand='海天',
                        company_spec='500ml',
                        company_unit='瓶',
                        company_category='调味品',
                        score=77.5,
                        status='candidate',
                        reasons=['名称接近'],
                    )
                ],
            )
            content = build_export_workbook(temp_path, [result])
            loaded = load_workbook(BytesIO(content))
            sheet = loaded.active
            self.assertEqual(sheet.cell(row=1, column=3).value, '匹配状态')
            self.assertEqual(sheet.cell(row=2, column=3).value, '建议匹配')
            self.assertEqual(sheet.cell(row=2, column=5).value, 'C001')
            self.assertEqual(sheet.cell(row=2, column=6).value, '海天酱油500ml')
        finally:
            temp_path.unlink(missing_ok=True)

    def test_manual_review_does_not_write_company_result_columns(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["商品名称", "编号"])
        sheet.append(["素鸡", "001"])

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            temp_path = Path(tmp.name)
        try:
            workbook.save(temp_path)
            result = MatchResult(
                customer_row_no=2,
                customer_name='素鸡',
                customer_brand='',
                customer_spec='',
                customer_unit='',
                customer_category='豆制品',
                match_status='manual_review',
                top_score=56.0,
                summary='名称接近，但需要人工审核',
                candidates=[
                    MatchCandidate(
                        company_row_no=9,
                        company_code='X001',
                        company_name='冻鸡骨架',
                        company_brand='',
                        company_spec='',
                        company_unit='袋',
                        company_category='冻品',
                        score=56.0,
                        status='candidate',
                        reasons=['名称接近'],
                    )
                ],
            )
            content = build_export_workbook(temp_path, [result])
            loaded = load_workbook(BytesIO(content))
            sheet = loaded.active
            self.assertEqual(sheet.cell(row=2, column=3).value, '待人工审核')
            self.assertEqual(sheet.cell(row=2, column=5).value, None)
            self.assertEqual(sheet.cell(row=2, column=6).value, None)
            self.assertEqual(sheet.cell(row=2, column=12).value, '待人工审核')
        finally:
            temp_path.unlink(missing_ok=True)

    def test_suggested_filename(self) -> None:
        self.assertEqual(suggested_export_name('客户商品库.xlsx'), '客户商品库_匹配结果.xlsx')


if __name__ == '__main__':
    unittest.main()
