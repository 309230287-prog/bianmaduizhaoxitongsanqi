from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from product_matcher.models import NormalizedRecord
from product_matcher.paths import DICTIONARY_DIR


DEFAULT_BRANDS: set[str] = set()
BRAND_DICTIONARY_FILE = DICTIONARY_DIR / "brands.txt"
PRODUCT_KEYWORD_FILE = DICTIONARY_DIR / "product_keywords.txt"

SUPPLY_CODE_RE = re.compile(r"[-－](?:CC|KK|DD|BB|AA)\d+$", re.IGNORECASE)
SPEC_RE = re.compile(
    r"(\[?\d+(?:\.\d+)?(?:\s*[*xX＊]\s*\d+(?:\.\d+)?){0,3}\s*(?:ml|mL|ML|l|L|kg|KG|g|G|斤|两|个|只|袋|包|盒|瓶|桶|箱|板|支|条)(?:/只|/个)?\]?)",
    re.IGNORECASE,
)
UNIT_RE = re.compile(r"(斤|个|件|包|袋|盒|瓶|桶|板|条|箱|只|杯|支|套|扎|块|卷|串)")


def build_normalized_samples(
    workbook_path: Path,
    mapping: dict[str, str],
    source_type: str,
    brand_dictionary: set[str] | None = None,
    product_dictionary: dict[str, str] | None = None,
    limit: int = 8,
) -> list[NormalizedRecord]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        records: list[NormalizedRecord] = []
        brand_dictionary = set(brand_dictionary or DEFAULT_BRANDS)
        product_dictionary = dict(product_dictionary or {})

        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            record = normalize_row(row_index, row, mapping, source_type, brand_dictionary, product_dictionary)
            if record.source_name:
                records.append(record)
            if len(records) >= limit:
                break
        return records
    finally:
        workbook.close()


def ensure_dictionary_files() -> None:
    DICTIONARY_DIR.mkdir(parents=True, exist_ok=True)
    for path in (BRAND_DICTIONARY_FILE, PRODUCT_KEYWORD_FILE):
        if not path.exists():
            path.write_text("", encoding="utf-8")


def load_text_dictionary(path: Path) -> set[str]:
    ensure_dictionary_files()
    terms: set[str] = set()
    for line in path.read_text(encoding="utf-8").splitlines():
        value = line.strip()
        if not value or value.startswith("#"):
            continue
        terms.add(value)
    return terms


def load_brand_dictionary(*sources: set[str]) -> set[str]:
    brand_dictionary = set(DEFAULT_BRANDS)
    brand_dictionary.update(load_text_dictionary(BRAND_DICTIONARY_FILE))
    for source in sources:
        for value in source:
            term = str(value or "").strip()
            if term:
                brand_dictionary.add(term)
    return brand_dictionary


def load_product_dictionary() -> dict[str, str]:
    ensure_dictionary_files()
    product_dictionary: dict[str, str] = {}
    for line in PRODUCT_KEYWORD_FILE.read_text(encoding="utf-8").splitlines():
        value = line.strip()
        if not value or value.startswith("#"):
            continue
        terms = [item.strip() for item in re.split(r"[|｜]", value) if item.strip()]
        if not terms:
            continue
        canonical = terms[0]
        for term in terms:
            lookup_key = _normalize_keyword_lookup(term)
            if lookup_key:
                product_dictionary[lookup_key] = canonical
    return product_dictionary


def collect_mapped_values(
    workbook_path: Path,
    mapping: dict[str, str],
    field_key: str,
    limit: int | None = None,
) -> set[str]:
    value_key = mapping.get(field_key)
    if not value_key:
        return set()

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        values: set[str] = set()
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            value = _get_mapped_value(row, value_key)
            if value:
                values.add(value)
            if limit and row_no >= limit:
                break
        return values
    finally:
        workbook.close()


def normalize_row(
    row_no: int,
    row: tuple[object, ...],
    mapping: dict[str, str],
    source_type: str,
    brand_dictionary: set[str],
    product_dictionary: dict[str, str] | None = None,
) -> NormalizedRecord:
    source_code = _get_mapped_value(row, mapping.get("product_code"))
    source_name = _get_mapped_value(row, mapping.get("product_name"))
    source_brand = _get_mapped_value(row, mapping.get("brand"))
    source_spec = _get_mapped_value(row, mapping.get("spec"))
    source_unit = _get_mapped_value(row, mapping.get("unit"))
    source_category = _get_mapped_value(row, mapping.get("category"))

    cleaned_name = clean_name(source_name)
    parsed_spec, spec_note = extract_spec(source_spec, cleaned_name)
    parsed_unit, unit_note = extract_unit(source_unit, cleaned_name, parsed_spec)
    parsed_brand, brand_note = extract_brand(source_brand, cleaned_name, brand_dictionary)
    parsed_name, name_note = extract_product_name(cleaned_name, parsed_brand, parsed_spec, product_dictionary)
    parsed_category = source_category.strip()

    notes = [note for note in (spec_note, unit_note, brand_note, name_note) if note]

    return NormalizedRecord(
        source_type=source_type,
        row_no=row_no,
        source_code=source_code,
        source_name=source_name,
        source_brand=source_brand,
        source_spec=source_spec,
        source_unit=source_unit,
        source_category=source_category,
        cleaned_name=cleaned_name,
        parsed_brand=parsed_brand,
        parsed_name=parsed_name,
        parsed_spec=parsed_spec,
        parsed_unit=parsed_unit,
        parsed_category=parsed_category,
        parse_notes="；".join(notes),
    )


def clean_name(value: str) -> str:
    text = str(value or "").strip()
    text = text.replace("（", "(").replace("）", ")")
    text = text.replace("＊", "*").replace("×", "*").replace("X", "x")
    text = re.sub(r"\s+", " ", text)
    text = SUPPLY_CODE_RE.sub("", text).strip("-_/ ")
    return text


def extract_spec(explicit_spec: str, cleaned_name: str) -> tuple[str, str]:
    if explicit_spec.strip():
        return normalize_spec(explicit_spec), "规格来自映射列"

    match = SPEC_RE.search(cleaned_name)
    if match:
        return normalize_spec(match.group(1)), "规格来自商品名称"

    return "", "规格未识别"


def normalize_spec(value: str) -> str:
    spec = str(value or "").strip().strip("[]")
    spec = spec.replace("＊", "*").replace("×", "*").replace("X", "x")
    spec = re.sub(r"\s+", "", spec)
    return spec


def extract_unit(explicit_unit: str, cleaned_name: str, parsed_spec: str) -> tuple[str, str]:
    if explicit_unit.strip():
        return explicit_unit.strip(), "单位来自映射列"

    for text, note in ((parsed_spec, "单位来自规格"), (cleaned_name, "单位来自商品名称")):
        match = UNIT_RE.search(text)
        if match:
            return match.group(1), note

    return "", "单位未识别"


def extract_brand(explicit_brand: str, cleaned_name: str, brand_dictionary: set[str]) -> tuple[str, str]:
    if explicit_brand.strip():
        return explicit_brand.strip(), "品牌来自映射列"

    for brand in sorted(brand_dictionary, key=len, reverse=True):
        if cleaned_name.startswith(brand) or f"{brand}牌" in cleaned_name:
            return brand, "品牌来自品牌词库"

    return "", "品牌未识别"


def extract_product_name(
    cleaned_name: str,
    parsed_brand: str,
    parsed_spec: str,
    product_dictionary: dict[str, str] | None = None,
) -> tuple[str, str]:
    product_name = cleaned_name
    if parsed_brand:
        product_name = product_name.replace(parsed_brand, " ", 1)
    if parsed_spec:
        product_name = product_name.replace(parsed_spec, " ")

    product_name = re.sub(r"[-_/()\[\]]", " ", product_name)
    product_name = re.sub(r"\s+", " ", product_name).strip()

    matched_product = _match_product_keyword(product_name, product_dictionary or {})
    if not matched_product and product_name != cleaned_name:
        matched_product = _match_product_keyword(cleaned_name, product_dictionary or {})
    if matched_product:
        return matched_product, "商品名称由商品词库归一"

    if product_name:
        return product_name, "商品名称由剩余核心词生成"
    return cleaned_name, "商品名称退回清洗后原文"


def _match_product_keyword(value: str, product_dictionary: dict[str, str]) -> str:
    lookup_value = _normalize_keyword_lookup(value)
    if not lookup_value:
        return ""

    matched_alias = ""
    matched_product = ""
    for alias, canonical in product_dictionary.items():
        if len(alias) < 2:
            continue
        if alias in lookup_value and len(alias) > len(matched_alias):
            matched_alias = alias
            matched_product = canonical
    return matched_product


def _normalize_keyword_lookup(value: str) -> str:
    normalized = clean_name(value).lower()
    return re.sub(r"[\s\-_/()\[\]]+", "", normalized)


def _get_mapped_value(row: tuple[object, ...], value_key: str | None) -> str:
    if not value_key:
        return ""
    try:
        index = int(value_key.replace("col_", "")) - 1
    except ValueError:
        return ""
    if index < 0 or index >= len(row):
        return ""
    cell = row[index]
    return "" if cell is None else str(cell).strip()
