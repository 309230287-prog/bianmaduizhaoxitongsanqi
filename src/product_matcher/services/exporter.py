from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from product_matcher.models import MatchResult

STATUS_LABELS = {
    "auto_matched": "自动匹配",
    "suggested": "建议匹配",
    "manual_review": "待人工审核",
    "unmatched": "未匹配",
}

STATUS_FILLS = {
    "auto_matched": PatternFill(fill_type="solid", fgColor="DDEEDB"),
    "suggested": PatternFill(fill_type="solid", fgColor="FFF0C9"),
    "manual_review": PatternFill(fill_type="solid", fgColor="FCE3D6"),
    "unmatched": PatternFill(fill_type="solid", fgColor="F8DADA"),
}

APPEND_HEADERS = [
    "匹配状态",
    "匹配分数",
    "我司商品编码",
    "我司商品名称",
    "我司品牌",
    "我司规格",
    "我司单位",
    "我司分类",
    "匹配说明",
    "审核状态",
    "审核备注",
]


SAFE_EXPORT_STATUSES = {"auto_matched", "suggested"}


def build_export_workbook(customer_workbook_path: Path, results: list[MatchResult]) -> bytes:
    workbook = load_workbook(customer_workbook_path)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header_start_col = sheet.max_column + 1

        for offset, header in enumerate(APPEND_HEADERS):
            cell = sheet.cell(row=1, column=header_start_col + offset)
            cell.value = header

        result_map = {result.customer_row_no: result for result in results}
        for row_no in range(2, sheet.max_row + 1):
            result = result_map.get(row_no)
            _write_result_row(sheet, row_no, header_start_col, result)

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def suggested_export_name(original_name: str) -> str:
    stem = Path(original_name).stem or "客户商品库"
    return f"{stem}_匹配结果.xlsx"


def _write_result_row(sheet, row_no: int, start_col: int, result: MatchResult | None) -> None:
    if result is None:
        values = ["未处理", "", "", "", "", "", "", "", "未生成匹配结果", "待审核", ""]
        fill = PatternFill(fill_type="solid", fgColor="EFEFEF")
    else:
        top = result.candidates[0] if result.candidates and result.match_status in SAFE_EXPORT_STATUSES else None
        review_status = "待人工审核" if result.match_status in {"manual_review", "unmatched"} else "待审核"
        values = [
            STATUS_LABELS.get(result.match_status, result.match_status),
            result.top_score,
            top.company_code if top else "",
            top.company_name if top else "",
            top.company_brand if top else "",
            top.company_spec if top else "",
            top.company_unit if top else "",
            top.company_category if top else "",
            result.summary,
            review_status,
            "",
        ]
        fill = STATUS_FILLS.get(result.match_status, PatternFill(fill_type="solid", fgColor="EFEFEF"))

    for offset, value in enumerate(values):
        cell = sheet.cell(row=row_no, column=start_col + offset)
        cell.value = value
        cell.fill = fill
