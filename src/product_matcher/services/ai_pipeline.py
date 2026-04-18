from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor, as_completed
import hashlib
import json
from dataclasses import asdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from product_matcher.models import MatchCandidate, MatchResult, NormalizedRecord
from product_matcher.services.ai_client import AIClientError, chat_json
from product_matcher.services import matching as matching_service
from product_matcher.services.normalization import normalize_row


NORMALIZE_BATCH_SIZE = 24
JUDGE_BATCH_SIZE = 20
RECALL_CANDIDATE_LIMIT = 3
MAX_PARALLEL_REQUESTS = 4


def build_ai_normalized_samples(
    workbook_path: Path,
    mapping: dict[str, str],
    source_type: str,
    runtime_settings: dict[str, Any],
    brand_dictionary: set[str],
    product_dictionary: dict[str, str],
    *,
    limit: int = 8,
    cache_dir: Path | None = None,
) -> list[NormalizedRecord]:
    cache_path = _normalized_cache_path(cache_dir, source_type, mapping, runtime_settings, limit)
    if cache_path and cache_path.exists():
        return _load_normalized_records(cache_path)

    fallback_records = _load_rule_normalized_records(
        workbook_path,
        mapping,
        source_type,
        brand_dictionary,
        product_dictionary,
        limit=limit,
    )
    enhanced_records = enhance_records_with_ai(fallback_records, runtime_settings)

    if cache_path:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(
            json.dumps([asdict(record) for record in enhanced_records], ensure_ascii=False),
            encoding="utf-8",
        )
    return enhanced_records


def enhance_records_with_ai(
    records: list[NormalizedRecord],
    runtime_settings: dict[str, Any],
) -> list[NormalizedRecord]:
    if not records:
        return []

    batches = [
        (batch_start, records[batch_start : batch_start + NORMALIZE_BATCH_SIZE])
        for batch_start in range(0, len(records), NORMALIZE_BATCH_SIZE)
    ]
    enhanced_batches: dict[int, list[NormalizedRecord]] = {}
    worker_count = min(MAX_PARALLEL_REQUESTS, len(batches))

    with ThreadPoolExecutor(max_workers=max(worker_count, 1)) as executor:
        future_map = {
            executor.submit(_normalize_batch, batch, runtime_settings): (batch_start, batch)
            for batch_start, batch in batches
        }
        for future in as_completed(future_map):
            batch_start, batch = future_map[future]
            try:
                ai_records = future.result()
                ai_map = {item["row_no"]: item for item in ai_records if isinstance(item.get("row_no"), int)}
                enhanced_batches[batch_start] = [
                    _merge_normalized_record(fallback, ai_map.get(fallback.row_no))
                    for fallback in batch
                ]
            except Exception:
                enhanced_batches[batch_start] = list(batch)

    enhanced: list[NormalizedRecord] = []
    for batch_start, _batch in batches:
        enhanced.extend(enhanced_batches[batch_start])
    return enhanced


def build_ai_match_preview(
    customer_records: list[NormalizedRecord],
    company_records: list[NormalizedRecord],
    runtime_settings: dict[str, Any],
    *,
    candidate_limit: int = 3,
    progress_callback=None,
) -> list[MatchResult]:
    index = matching_service._build_company_index(company_records)
    queued: list[tuple[NormalizedRecord, list[MatchCandidate]]] = []
    results: list[MatchResult] = []

    for customer in customer_records:
        candidates = matching_service._score_candidates(customer, company_records, index, RECALL_CANDIDATE_LIMIT)
        if not candidates:
            results.append(
                MatchResult(
                    customer_row_no=customer.row_no,
                    customer_name=customer.source_name,
                    customer_brand=customer.parsed_brand,
                    customer_spec=customer.parsed_spec,
                    customer_unit=customer.parsed_unit,
                    customer_category=customer.parsed_category,
                    match_status="unmatched",
                    top_score=0.0,
                    summary="没有召回到合理候选",
                    candidates=[],
                )
            )
            continue
        queued.append((customer, candidates))

    batches = [
        (batch_start, queued[batch_start : batch_start + JUDGE_BATCH_SIZE])
        for batch_start in range(0, len(queued), JUDGE_BATCH_SIZE)
    ]
    judged_batches: dict[int, list[MatchResult]] = {}
    worker_count = min(MAX_PARALLEL_REQUESTS, len(batches))

    with ThreadPoolExecutor(max_workers=max(worker_count, 1)) as executor:
        future_map = {
            executor.submit(_judge_customer_batch, batch, runtime_settings): (batch_start, batch)
            for batch_start, batch in batches
        }
        completed_batches = 0
        total_batches = len(batches)
        for future in as_completed(future_map):
            batch_start, batch = future_map[future]
            try:
                ai_results = future.result()
                ai_result_map = {
                    item["customer_row_no"]: item
                    for item in ai_results
                    if isinstance(item, dict) and isinstance(item.get("customer_row_no"), int)
                }
                judged_batches[batch_start] = [
                    _build_ai_match_result(
                        customer,
                        candidates,
                        ai_result_map.get(customer.row_no),
                        candidate_limit,
                    )
                    for customer, candidates in batch
                ]
            except Exception:
                judged_batches[batch_start] = [
                    _build_rule_result_from_candidates(customer, candidates, candidate_limit)
                    for customer, candidates in batch
                ]
            completed_batches += 1
            if progress_callback is not None:
                progress_callback(completed_batches, total_batches)

    for batch_start, _batch in batches:
        results.extend(judged_batches[batch_start])

    results.sort(key=lambda item: item.customer_row_no)
    return results


def _load_rule_normalized_records(
    workbook_path: Path,
    mapping: dict[str, str],
    source_type: str,
    brand_dictionary: set[str],
    product_dictionary: dict[str, str],
    *,
    limit: int,
) -> list[NormalizedRecord]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        records: list[NormalizedRecord] = []
        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            record = normalize_row(row_index, row, mapping, source_type, brand_dictionary, product_dictionary)
            if record.source_name:
                records.append(record)
            if len(records) >= limit:
                break
    finally:
        workbook.close()
    return records


def _normalize_batch(batch: list[NormalizedRecord], runtime_settings: dict[str, Any]) -> list[dict[str, Any]]:
    payload = {
        "task": "normalize_products",
        "items": [
            {
                "row_no": record.row_no,
                "source_type": record.source_type,
                "source_code": record.source_code,
                "source_name": record.source_name,
                "source_brand": record.source_brand,
                "source_spec": record.source_spec,
                "source_unit": record.source_unit,
                "source_category": record.source_category,
                "cleaned_name": record.cleaned_name,
                "parsed_brand": record.parsed_brand,
                "parsed_name": record.parsed_name,
                "parsed_spec": record.parsed_spec,
                "parsed_unit": record.parsed_unit,
                "parsed_category": record.parsed_category,
            }
            for record in batch
        ],
    }
    system_prompt = """
你是商品标准化助手。
请根据输入的原始商品信息，补充并归一以下字段：
- brand_std
- product_std
- spec_std
- package_std
- unit_std
- category_std
- alias_notes
- normalize_confidence

要求：
1. 只返回 JSON 对象，格式必须是 {"items":[...]}。
2. items 中每个元素必须包含 row_no。
3. normalize_confidence 使用 0 到 1 之间的小数。
4. 缺失字段可以返回空字符串，不要编造供应商编码。
5. 优先减少误识别，拿不准时保留原有语义。
""".strip()
    response = chat_json(runtime_settings, system_prompt, payload, max_tokens=1800)
    items = response.get("items", [])
    if not isinstance(items, list):
        raise AIClientError("标准化结果缺少 items 数组。")
    return [item for item in items if isinstance(item, dict)]


def _merge_normalized_record(record: NormalizedRecord, ai_item: dict[str, Any] | None) -> NormalizedRecord:
    if not ai_item:
        return record

    parsed_brand = _clean_text(ai_item.get("brand_std")) or record.parsed_brand
    parsed_name = _clean_text(ai_item.get("product_std")) or record.parsed_name or record.cleaned_name
    parsed_spec = _clean_text(ai_item.get("spec_std")) or record.parsed_spec
    parsed_unit = _clean_text(ai_item.get("unit_std")) or record.parsed_unit
    parsed_category = _clean_text(ai_item.get("category_std")) or record.parsed_category
    package_std = _clean_text(ai_item.get("package_std"))
    alias_notes = _clean_text(ai_item.get("alias_notes"))
    confidence = _clean_text(ai_item.get("normalize_confidence"))

    notes = ["模型标准化"]
    if package_std:
        notes.append(f"包装 {package_std}")
    if alias_notes:
        notes.append(alias_notes)
    if confidence:
        notes.append(f"置信度 {confidence}")
    if record.parse_notes:
        notes.append(record.parse_notes)

    return NormalizedRecord(
        source_type=record.source_type,
        row_no=record.row_no,
        source_code=record.source_code,
        source_name=record.source_name,
        source_brand=record.source_brand,
        source_spec=record.source_spec,
        source_unit=record.source_unit,
        source_category=record.source_category,
        cleaned_name=record.cleaned_name,
        parsed_brand=parsed_brand,
        parsed_name=parsed_name,
        parsed_spec=parsed_spec,
        parsed_unit=parsed_unit,
        parsed_category=parsed_category,
        parse_notes="；".join(item for item in notes if item),
    )


def _judge_customer_batch(
    batch: list[tuple[NormalizedRecord, list[MatchCandidate]]],
    runtime_settings: dict[str, Any],
) -> list[dict[str, Any]]:
    payload = {
        "task": "judge_product_matches",
        "items": [
            {
                "customer_row_no": customer.row_no,
                "customer": {
                    "source_name": customer.source_name,
                    "brand_std": customer.parsed_brand,
                    "product_std": customer.parsed_name,
                    "spec_std": customer.parsed_spec,
                    "unit_std": customer.parsed_unit,
                    "category_std": customer.parsed_category,
                    "parse_notes": customer.parse_notes,
                },
                "candidates": [
                    {
                        "company_row_no": candidate.company_row_no,
                        "company_code": candidate.company_code,
                        "company_name": candidate.company_name,
                        "brand_std": candidate.company_brand,
                        "spec_std": candidate.company_spec,
                        "unit_std": candidate.company_unit,
                        "category_std": candidate.company_category,
                        "rule_score": candidate.score,
                        "rule_status": candidate.status,
                        "rule_reasons": candidate.reasons,
                    }
                    for candidate in candidates
                ],
            }
            for customer, candidates in batch
        ],
    }
    system_prompt = """
你是商品匹配裁决助手。
请在每个 customer 的候选列表里，挑出最合适的我司商品；如果都不可靠，就返回 unmatched。

要求：
1. 只返回 JSON 对象，格式必须是 {"items":[...]}。
2. 每个 item 必须包含：
   - customer_row_no
   - match_status
   - selected_candidate_row_no
   - confidence
   - summary
   - reasons
   - candidate_rankings
3. match_status 只能是 auto_matched、suggested、manual_review、unmatched。
4. candidate_rankings 必须是按推荐顺序排列的 company_row_no 数组。
5. confidence 使用 0 到 1 之间的小数。
6. 目标是尽量少错配，不要激进自动通过。
7. 如果品牌、规格、单位存在明显硬冲突，优先给 manual_review 或 unmatched。
""".strip()
    response = chat_json(runtime_settings, system_prompt, payload, max_tokens=2200)
    items = response.get("items", [])
    if not isinstance(items, list):
        raise AIClientError("裁决结果缺少 items 数组。")
    return [item for item in items if isinstance(item, dict)]


def _build_ai_match_result(
    customer: NormalizedRecord,
    candidates: list[MatchCandidate],
    ai_item: dict[str, Any] | None,
    candidate_limit: int,
) -> MatchResult:
    if not ai_item:
        return _build_rule_result_from_candidates(customer, candidates, candidate_limit)

    ranking_rows = ai_item.get("candidate_rankings", [])
    ordered_candidates = _reorder_candidates(candidates, ranking_rows)
    selected_row_no = _to_int(ai_item.get("selected_candidate_row_no"))
    selected_candidate = _find_selected_candidate(ordered_candidates, selected_row_no)

    match_status = _clean_text(ai_item.get("match_status")) or "manual_review"
    if match_status not in {"auto_matched", "suggested", "manual_review", "unmatched"}:
        match_status = "manual_review"

    confidence = _to_float(ai_item.get("confidence"))
    match_status = _apply_safety_guards(match_status, selected_candidate, confidence)
    summary = _clean_text(ai_item.get("summary")) or "模型已完成匹配裁决"
    reasons = _clean_reason_list(ai_item.get("reasons"))

    merged_candidates: list[MatchCandidate] = []
    for index, candidate in enumerate(ordered_candidates[:candidate_limit]):
        candidate_reasons = list(candidate.reasons)
        if selected_candidate and candidate.company_row_no == selected_candidate.company_row_no:
            candidate_reasons = reasons + candidate_reasons
        candidate_score = candidate.score
        if index == 0:
            candidate_score = max(candidate_score, round(confidence * 100, 2))
        merged_candidates.append(
            MatchCandidate(
                company_row_no=candidate.company_row_no,
                company_code=candidate.company_code,
                company_name=candidate.company_name,
                company_brand=candidate.company_brand,
                company_spec=candidate.company_spec,
                company_unit=candidate.company_unit,
                company_category=candidate.company_category,
                score=candidate_score,
                status=candidate.status,
                reasons=candidate_reasons,
            )
        )

    if match_status == "unmatched":
        merged_candidates = []

    top_score = merged_candidates[0].score if merged_candidates else 0.0
    return MatchResult(
        customer_row_no=customer.row_no,
        customer_name=customer.source_name,
        customer_brand=customer.parsed_brand,
        customer_spec=customer.parsed_spec,
        customer_unit=customer.parsed_unit,
        customer_category=customer.parsed_category,
        match_status=match_status,
        top_score=top_score,
        summary=summary,
        candidates=merged_candidates,
    )


def _build_rule_result_from_candidates(
    customer: NormalizedRecord,
    candidates: list[MatchCandidate],
    candidate_limit: int,
) -> MatchResult:
    limited_candidates = candidates[:candidate_limit]
    if not limited_candidates:
        return MatchResult(
            customer_row_no=customer.row_no,
            customer_name=customer.source_name,
            customer_brand=customer.parsed_brand,
            customer_spec=customer.parsed_spec,
            customer_unit=customer.parsed_unit,
            customer_category=customer.parsed_category,
            match_status="unmatched",
            top_score=0.0,
            summary="没有召回到合理候选",
            candidates=[],
        )

    top = limited_candidates[0]
    return MatchResult(
        customer_row_no=customer.row_no,
        customer_name=customer.source_name,
        customer_brand=customer.parsed_brand,
        customer_spec=customer.parsed_spec,
        customer_unit=customer.parsed_unit,
        customer_category=customer.parsed_category,
        match_status=matching_service._grade_result(top),
        top_score=top.score,
        summary="；".join(top.reasons[:3]) or "已生成候选",
        candidates=limited_candidates,
    )


def _apply_safety_guards(
    match_status: str,
    selected_candidate: MatchCandidate | None,
    confidence: float,
) -> str:
    if match_status == "unmatched":
        return "unmatched"
    if selected_candidate is None:
        return "manual_review"
    if selected_candidate.status == "hard_conflict":
        return "manual_review"
    if match_status == "auto_matched" and confidence < 0.85:
        return "suggested"
    if match_status == "suggested" and confidence < 0.55:
        return "manual_review"
    return match_status


def _find_selected_candidate(
    ordered_candidates: list[MatchCandidate],
    selected_row_no: int | None,
) -> MatchCandidate | None:
    if selected_row_no is not None:
        for candidate in ordered_candidates:
            if candidate.company_row_no == selected_row_no:
                return candidate
    if ordered_candidates:
        return ordered_candidates[0]
    return None


def _reorder_candidates(candidates: list[MatchCandidate], ranking_rows: Any) -> list[MatchCandidate]:
    ranking: list[int] = []
    if isinstance(ranking_rows, list):
        for item in ranking_rows:
            row_no = _to_int(item)
            if row_no is not None:
                ranking.append(row_no)

    ordered: list[MatchCandidate] = []
    seen: set[int] = set()
    for row_no in ranking:
        for candidate in candidates:
            if candidate.company_row_no == row_no and row_no not in seen:
                ordered.append(candidate)
                seen.add(row_no)
                break
    for candidate in candidates:
        if candidate.company_row_no not in seen:
            ordered.append(candidate)
    return ordered


def _normalized_cache_path(
    cache_dir: Path | None,
    source_type: str,
    mapping: dict[str, str],
    runtime_settings: dict[str, Any],
    limit: int,
) -> Path | None:
    if cache_dir is None:
        return None
    fingerprint = hashlib.sha1(
        json.dumps(
            {
                "mapping": mapping,
                "model": runtime_settings.get("production_model_name"),
                "limit": limit,
            },
            ensure_ascii=False,
            sort_keys=True,
        ).encode("utf-8")
    ).hexdigest()[:12]
    return cache_dir / f"{source_type}_normalized_ai_{fingerprint}.json"


def _load_normalized_records(path: Path) -> list[NormalizedRecord]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    records: list[NormalizedRecord] = []
    for item in payload:
        if isinstance(item, dict):
            records.append(NormalizedRecord(**item))
    return records


def _clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _clean_reason_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = _clean_text(value)
    return [text] if text else []


def _to_int(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _to_float(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0
