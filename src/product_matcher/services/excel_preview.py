from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from product_matcher.models import ColumnOption, WorkbookPreview


class WorkbookPreviewError(ValueError):
    """Raised when an uploaded workbook cannot be previewed safely."""


SUPPORTED_EXTENSIONS = {".xlsx"}
MAX_SAMPLE_ROWS = 8

FIELD_HINTS = {
    "company": {
        "product_code": ["spuid", "商品编码", "编号", "自定义编码", "编码"],
        "product_name": ["spu名称", "商品名称", "名称"],
        "brand": ["品牌", "厂牌", "牌子"],
        "spec": ["规格", "描述"],
        "unit": ["基本单位", "单位"],
        "category": ["二级分类名称", "三级分类名称", "分类名称", "分类"],
    },
    "customer": {
        "product_code": ["商品编码", "编号", "编码"],
        "product_name": ["商品名称", "名称", "产品名称"],
        "brand": ["品牌", "厂牌", "牌子"],
        "spec": ["规格", "描述"],
        "unit": ["单位"],
        "category": ["类别", "分类", "大类"],
    },
}


def build_workbook_preview(filename: str, content: bytes) -> WorkbookPreview:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise WorkbookPreviewError("第一版目前只支持 .xlsx 文件。")

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover
        raise WorkbookPreviewError(f"无法读取 Excel 文件：{exc}") from exc

    try:
        sheet_name = workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, tuple())
        columns = _build_columns(header_row)
        sample_rows = _build_sample_rows(rows, len(columns))
        total_rows = max((sheet.max_row or 1) - 1, 0)

        return WorkbookPreview(
            filename=filename,
            sheet_name=sheet_name,
            total_rows=total_rows,
            columns=columns,
            sample_rows=sample_rows,
        )
    finally:
        workbook.close()


def suggest_mapping(columns: list[ColumnOption], source_type: str) -> dict[str, str]:
    hints = FIELD_HINTS[source_type]
    suggestions: dict[str, str] = {}
    used_values: set[str] = set()

    for field_key, keywords in hints.items():
        suggestion = _find_best_column(columns, keywords, used_values)
        if suggestion:
            suggestions[field_key] = suggestion.value_key
            used_values.add(suggestion.value_key)
    return sanitize_mapping(columns, source_type, suggestions)


def sanitize_mapping(columns: list[ColumnOption], source_type: str, mapping: dict[str, str]) -> dict[str, str]:
    hints = FIELD_HINTS[source_type]
    by_value = {column.value_key: column for column in columns}
    cleaned: dict[str, str] = {}

    for field_key in hints:
        value_key = mapping.get(field_key, "")
        column = by_value.get(value_key)
        if not column:
            cleaned[field_key] = ""
            continue
        header = column.raw_header.lower()
        if any(keyword.lower() in header for keyword in hints[field_key]):
            cleaned[field_key] = value_key
        else:
            cleaned[field_key] = ""
    return cleaned


def describe_mapping(columns: list[ColumnOption], mapping: dict[str, str]) -> dict[str, dict[str, str]]:
    by_value = {column.value_key: column for column in columns}
    payload: dict[str, dict[str, str]] = {}
    for field_key, value_key in mapping.items():
        column = by_value.get(value_key)
        payload[field_key] = {
            "value_key": value_key,
            "header": column.raw_header if column else "",
            "label": column.label if column else "",
        }
    return payload


def _find_best_column(columns: list[ColumnOption], keywords: list[str], used_values: set[str]) -> ColumnOption | None:
    for keyword in keywords:
        for column in columns:
            if column.value_key in used_values:
                continue
            if keyword.lower() in column.raw_header.lower():
                return column
    return None


def _build_columns(header_row: tuple[object, ...]) -> list[ColumnOption]:
    headers = [_normalize_header(cell) for cell in header_row]
    columns: list[ColumnOption] = []
    occurrences: dict[str, int] = {}
    totals: dict[str, int] = {}

    for header in headers:
        totals[header] = totals.get(header, 0) + 1

    for index, header in enumerate(headers, start=1):
        occurrences[header] = occurrences.get(header, 0) + 1
        display = header
        if totals[header] > 1:
            display = f"{header}（重复列{occurrences[header]}）"
        columns.append(ColumnOption(index=index, raw_header=header, label=display, value_key=f"col_{index}"))
    return columns


def _normalize_header(value: object) -> str:
    if value is None:
        return "空白列"
    text = str(value).strip()
    return text or "空白列"


def _build_sample_rows(rows, column_count: int) -> list[list[str]]:
    sample_rows: list[list[str]] = []
    for row in rows:
        normalized = ["" if cell is None else str(cell).strip() for cell in row[:column_count]]
        if any(normalized):
            sample_rows.append(normalized)
        if len(sample_rows) >= MAX_SAMPLE_ROWS:
            break
    return sample_rows
