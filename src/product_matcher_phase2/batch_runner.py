from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from product_matcher_phase2.candidate_generation import generate_candidates
from product_matcher_phase2.model_io import build_model_input_payload
from product_matcher_phase2.model_trial import ModelCaller, ModelTrialCase, ModelTrialResult, run_trial_cases
from product_matcher_phase2.schemas import CompanyProduct, CustomerRecord, ResultStatus


@dataclass(frozen=True)
class Phase2BatchRow:
    customer_record: CustomerRecord
    result: ModelTrialResult
    selected_product_name: str
    top_candidate_code: str
    top_candidate_name: str


def build_phase2_batch_cases(
    customer_records: Iterable[CustomerRecord],
    company_products: Iterable[CompanyProduct],
    *,
    candidate_limit: int = 10,
    max_rows: int | None = None,
) -> list[ModelTrialCase]:
    products = list(company_products)
    cases: list[ModelTrialCase] = []
    for record in _take(customer_records, max_rows):
        candidates = generate_candidates(record, products, limit=candidate_limit)
        cases.append(
            ModelTrialCase(
                sample_id=record.record_id,
                sample_group="live_batch",
                expected_company_code="",
                expected_result_status="",
                payload=build_model_input_payload(record, candidates, applicable_memories=[]),
            )
        )
    return cases


def run_phase2_batch(
    customer_records: Iterable[CustomerRecord],
    company_products: Iterable[CompanyProduct],
    call_model: ModelCaller,
    *,
    candidate_limit: int = 10,
    max_rows: int | None = None,
) -> list[Phase2BatchRow]:
    records = list(_take(customer_records, max_rows))
    cases = build_phase2_batch_cases(
        records,
        company_products,
        candidate_limit=candidate_limit,
        max_rows=None,
    )
    results = run_trial_cases(cases, call_model)
    return [
        _build_batch_row(record, case, result)
        for record, case, result in zip(records, cases, results, strict=True)
    ]


def summarize_batch_results(rows: Iterable[Phase2BatchRow]) -> dict[str, Any]:
    row_list = list(rows)
    total_count = len(row_list)
    json_valid_count = sum(1 for row in row_list if row.result.parse_ok)
    auto_code_count = sum(1 for row in row_list if row.result.can_auto_code)
    suggested_or_review_count = sum(
        1
        for row in row_list
        if row.result.parsed_result_status
        in {ResultStatus.SUGGESTED_CODE.value, ResultStatus.MANUAL_REVIEW.value}
    )
    unmatched_count = sum(
        1 for row in row_list if row.result.parsed_result_status == ResultStatus.UNMATCHED.value
    )
    model_error_count = sum(
        1 for row in row_list if row.result.parsed_result_status == ResultStatus.MODEL_ERROR.value
    )
    return {
        "total_count": total_count,
        "json_valid_count": json_valid_count,
        "json_valid_rate": json_valid_count / total_count if total_count else 0.0,
        "auto_code_count": auto_code_count,
        "suggested_or_review_count": suggested_or_review_count,
        "unmatched_count": unmatched_count,
        "model_error_count": model_error_count,
    }


PHASE2_BATCH_HEADERS = [
    "客户行号",
    "客户记录编号",
    "客户编码",
    "客户商品名称",
    "客户规格",
    "客户单位",
    "客户品牌",
    "处理结论",
    "能否自动落码",
    "对照我司编码",
    "对照我司商品名称",
    "首候选我司编码",
    "首候选商品名称",
    "人工复核说明",
    "匹配依据",
    "异常说明",
]


def write_phase2_batch_results_xlsx(rows: Iterable[Phase2BatchRow], path: str | Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "编码对照表"
    worksheet.append(PHASE2_BATCH_HEADERS)
    for row in rows:
        mapped = row.customer_record.mapped_fields
        result = row.result
        worksheet.append(
            [
                row.customer_record.source_row_number,
                row.customer_record.record_id,
                mapped.get("code", ""),
                mapped.get("name", ""),
                mapped.get("spec", ""),
                mapped.get("unit", ""),
                mapped.get("brand", ""),
                _human_status(result.parsed_result_status),
                "是，可自动落码" if result.can_auto_code else "否，需人工复核",
                result.selected_company_code,
                row.selected_product_name,
                row.top_candidate_code,
                row.top_candidate_name,
                result.manual_review_reason,
                result.evidence_summary,
                result.error_message,
            ]
        )
    _style_code_mapping_sheet(worksheet)

    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def _human_status(status: str) -> str:
    return {
        ResultStatus.STRONG_AUTO_CODE.value: "可自动落码",
        ResultStatus.WEAK_AUTO_CODE.value: "可自动落码",
        ResultStatus.SUGGESTED_CODE.value: "建议编码",
        ResultStatus.MANUAL_REVIEW.value: "需人工审核",
        ResultStatus.UNMATCHED.value: "未匹配",
        ResultStatus.MODEL_ERROR.value: "模型异常",
    }.get(status, "需人工审核")


def _style_code_mapping_sheet(worksheet) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="F4E6D4")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    widths = {
        "A": 10,
        "B": 18,
        "C": 16,
        "D": 32,
        "E": 18,
        "F": 10,
        "G": 14,
        "H": 14,
        "I": 16,
        "J": 18,
        "K": 34,
        "L": 18,
        "M": 34,
        "N": 36,
        "O": 40,
        "P": 28,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A2"


def _build_batch_row(
    record: CustomerRecord,
    case: ModelTrialCase,
    result: ModelTrialResult,
) -> Phase2BatchRow:
    candidates = case.payload.get("candidate_products") or []
    selected_product = _find_candidate_product(candidates, result.selected_candidate_id)
    top_product = candidates[0].get("product", {}) if candidates else {}
    return Phase2BatchRow(
        customer_record=record,
        result=result,
        selected_product_name=str(selected_product.get("name", "")),
        top_candidate_code=str(top_product.get("code", "")),
        top_candidate_name=str(top_product.get("name", "")),
    )


def _find_candidate_product(candidates: list[dict[str, Any]], candidate_id: str) -> dict[str, Any]:
    if not candidate_id:
        return {}
    for candidate in candidates:
        if candidate.get("candidate_id") == candidate_id:
            product = candidate.get("product") or {}
            return product if isinstance(product, dict) else {}
    return {}


def _take(records: Iterable[CustomerRecord], max_rows: int | None) -> list[CustomerRecord]:
    record_list: list[CustomerRecord] = []
    for record in records:
        if max_rows is not None and len(record_list) >= max_rows:
            break
        record_list.append(record)
    return record_list
