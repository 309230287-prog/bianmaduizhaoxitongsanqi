from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from product_matcher_phase2.schemas import CompanyProduct, CustomerRecord


def _stringify(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _unique_headers(raw_headers: Iterable[object]) -> list[str]:
    counts: dict[str, int] = {}
    headers: list[str] = []
    for index, value in enumerate(raw_headers, start=1):
        header = _stringify(value) or f"未命名列{index}"
        counts[header] = counts.get(header, 0) + 1
        if counts[header] == 1:
            headers.append(header)
        else:
            headers.append(f"{header}_{counts[header]}")
    return headers


def _load_rows(
    path: str | Path,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> list[tuple[int, dict[str, str]]]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    header_values = next(
        worksheet.iter_rows(
            min_row=header_row,
            max_row=header_row,
            values_only=True,
        )
    )
    headers = _unique_headers(header_values)
    rows: list[tuple[int, dict[str, str]]] = []
    for row_number, row_values in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        raw_fields = {
            header: _stringify(value)
            for header, value in zip(headers, row_values, strict=False)
        }
        if any(raw_fields.values()):
            rows.append((row_number, raw_fields))
    workbook.close()
    return rows


def _value_by_header_base(
    raw_fields: dict[str, str],
    header_bases: Iterable[str],
    prefer_last: bool = False,
) -> str:
    for header_base in header_bases:
        matches = [
            value
            for header, value in raw_fields.items()
            if header == header_base or header.startswith(f"{header_base}_")
        ]
        if prefer_last:
            matches = list(reversed(matches))
        for value in matches:
            if value:
                return value
    return ""


def _compact_mapping(mapping: dict[str, str]) -> dict[str, str]:
    return {key: value for key, value in mapping.items() if value}


def load_customer_records(
    path: str | Path,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> list[CustomerRecord]:
    records: list[CustomerRecord] = []
    for row_number, raw_fields in _load_rows(path, sheet_name=sheet_name, header_row=header_row):
        mapped_fields = _compact_mapping(
            {
                "code": _value_by_header_base(raw_fields, ["编号", "商品编码", "编码"]),
                "category": _value_by_header_base(raw_fields, ["类别", "分类"]),
                "name": _value_by_header_base(raw_fields, ["商品名称", "品名", "名称"], prefer_last=True),
                "spec": _value_by_header_base(raw_fields, ["规格", "规格型号"]),
                "unit": _value_by_header_base(raw_fields, ["单位", "基本单位"]),
                "brand": _value_by_header_base(raw_fields, ["品牌"]),
                "remark": _value_by_header_base(raw_fields, ["备注", "说明", "未命名列9"]),
            }
        )
        records.append(
            CustomerRecord(
                record_id=f"customer:{row_number}",
                source_row_number=row_number,
                raw_fields=raw_fields,
                mapped_fields=mapped_fields,
            )
        )
    return records


def load_company_products(
    path: str | Path,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> list[CompanyProduct]:
    products: list[CompanyProduct] = []
    for row_number, raw_fields in _load_rows(path, sheet_name=sheet_name, header_row=header_row):
        product_id = _value_by_header_base(raw_fields, ["SPUID", "商品编码", "编码"]) or f"company:{row_number}"
        category_path = [
            value
            for value in [
                _value_by_header_base(raw_fields, ["一级分类名称"]),
                _value_by_header_base(raw_fields, ["二级分类名称"]),
                _value_by_header_base(raw_fields, ["三级分类名称"]),
            ]
            if value
        ]
        products.append(
            CompanyProduct(
                product_id=product_id,
                code=product_id,
                name=_value_by_header_base(raw_fields, ["SPU名称（可修改）", "SPU名称", "商品名称", "名称"]),
                unit=_value_by_header_base(raw_fields, ["SPU基本单位", "单位", "基本单位"]),
                alias=_value_by_header_base(raw_fields, ["SPU别名（可修改）", "别名"]),
                description=_value_by_header_base(raw_fields, ["SPU描述（可修改）", "描述", "说明"]),
                category_path=category_path,
                raw_fields=raw_fields,
            )
        )
    return products
