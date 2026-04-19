from __future__ import annotations

import json
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from pydantic import ValidationError

from product_matcher_phase2.schemas import ModelDecision, normalize_model_decision_payload


DIAGNOSTIC_CATEGORIES = (
    "invalid_json",
    "schema_validation_error",
    "business_rule_validation_error",
    "model_call_error",
    "empty_output",
    "schema_valid_after_normalization",
    "unknown",
)


@dataclass(frozen=True)
class TrialResultRow:
    sample_id: str
    sample_group: str
    expected_company_code: str
    expected_result_status: str
    parsed_result_status: str
    status_matches_expected: Any
    parse_ok: Any
    selected_candidate_id: str
    selected_company_code: str
    can_auto_code: Any
    evidence_summary: str
    manual_review_reason: str
    error_message: str
    raw_model_output: str


def load_trial_result_rows(path: str | Path) -> list[TrialResultRow]:
    workbook = load_workbook(path, read_only=True)
    try:
        worksheet = workbook.active
        headers = [str(cell.value or "") for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        header_index = {name: index for index, name in enumerate(headers)}
        rows: list[TrialResultRow] = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            row = {header: values[index] if index < len(values) else "" for header, index in header_index.items()}
            rows.append(
                TrialResultRow(
                    sample_id=str(row.get("sample_id", "") or ""),
                    sample_group=str(row.get("sample_group", "") or ""),
                    expected_company_code=str(row.get("expected_company_code", "") or ""),
                    expected_result_status=str(row.get("expected_result_status", "") or ""),
                    parsed_result_status=str(row.get("parsed_result_status", "") or ""),
                    status_matches_expected=row.get("status_matches_expected", ""),
                    parse_ok=row.get("parse_ok", ""),
                    selected_candidate_id=str(row.get("selected_candidate_id", "") or ""),
                    selected_company_code=str(row.get("selected_company_code", "") or ""),
                    can_auto_code=row.get("can_auto_code", ""),
                    evidence_summary=str(row.get("evidence_summary", "") or ""),
                    manual_review_reason=str(row.get("manual_review_reason", "") or ""),
                    error_message=str(row.get("error_message", "") or ""),
                    raw_model_output=str(row.get("raw_model_output", "") or ""),
                )
            )
        return rows
    finally:
        workbook.close()


def diagnose_trial_results_xlsx(path: str | Path) -> dict[str, Any]:
    rows = load_trial_result_rows(path)
    categories: dict[str, list[dict[str, Any]]] = {category: [] for category in DIAGNOSTIC_CATEGORIES}
    category_counts: Counter[str] = Counter()
    for row in rows:
        category, reason = classify_trial_failure(row)
        category_counts[category] += 1
        if len(categories[category]) < 3:
            categories[category].append(
                {
                    "sample_id": row.sample_id,
                    "sample_group": row.sample_group,
                    "reason": reason,
                    "parsed_result_status": row.parsed_result_status,
                    "error_message": _shorten(row.error_message, 180),
                }
            )

    return {
        "source_path": str(Path(path)),
        "total_count": len(rows),
        "category_counts": {category: category_counts.get(category, 0) for category in DIAGNOSTIC_CATEGORIES},
        "representative_samples": categories,
        "acceptance_summary": summarize_acceptance(rows),
    }


def summarize_acceptance(rows: Iterable[TrialResultRow]) -> dict[str, int]:
    row_list = list(rows)
    selected_code_match_count = sum(1 for row in row_list if _selected_code_matches_expected(row))
    selected_code_mismatch_count = sum(1 for row in row_list if _selected_code_mismatches_expected(row))
    auto_code_count = sum(1 for row in row_list if _row_can_auto_code(row))
    unsafe_auto_code_count = sum(1 for row in row_list if _is_unsafe_auto_code(row))
    return {
        "status_match_count": sum(1 for row in row_list if _to_bool(row.status_matches_expected)),
        "selected_code_match_count": selected_code_match_count,
        "selected_code_mismatch_count": selected_code_mismatch_count,
        "auto_code_count": auto_code_count,
        "unsafe_auto_code_count": unsafe_auto_code_count,
    }


def classify_trial_failure(row: TrialResultRow) -> tuple[str, str]:
    raw_output = (row.raw_model_output or "").strip()
    error_message = (row.error_message or "").strip()

    if not raw_output:
        if error_message:
            return "model_call_error", _shorten(error_message)
        return "empty_output", "模型返回空内容"

    try:
        payload = json.loads(row.raw_model_output)
    except json.JSONDecodeError as exc:
        return "invalid_json", f"JSON 解析失败：{_shorten(str(exc))}"

    try:
        ModelDecision.model_validate(normalize_model_decision_payload(payload))
    except ValidationError as exc:
        errors = exc.errors()
        if _has_business_rule_error(errors):
            return "business_rule_validation_error", _business_rule_reason(errors)
        if _has_schema_error(errors):
            return "schema_validation_error", _schema_reason(errors)
        return "unknown", _shorten(str(exc), 240)

    return "schema_valid_after_normalization", "历史结果表标记失败，但原始输出经当前风险标记规范化后已通过模型校验"


def render_trial_diagnostics_markdown(summary: dict[str, Any]) -> str:
    lines = [
        "# Phase 2 Trial Diagnostics",
        "",
        f"- Source: `{summary.get('source_path', '')}`",
        f"- Total rows: {summary.get('total_count', 0)}",
        "",
        "## Category Counts",
    ]
    for category in DIAGNOSTIC_CATEGORIES:
        lines.append(f"- `{category}`: {summary.get('category_counts', {}).get(category, 0)}")

    lines.extend(["", "## Acceptance Summary"])
    acceptance_summary = summary.get("acceptance_summary", {})
    for key in (
        "status_match_count",
        "selected_code_match_count",
        "selected_code_mismatch_count",
        "auto_code_count",
        "unsafe_auto_code_count",
    ):
        lines.append(f"- `{key}`: {acceptance_summary.get(key, 0)}")

    lines.extend(["", "## Representative Samples"])
    representative_samples = summary.get("representative_samples", {})
    for category in DIAGNOSTIC_CATEGORIES:
        lines.append(f"### {category}")
        samples = representative_samples.get(category, [])
        if not samples:
            lines.append("- None")
            lines.append("")
            continue
        for sample in samples:
            lines.append(
                f"- `{sample.get('sample_id', '')}` ({sample.get('sample_group', '')}): "
                f"{sample.get('reason', '')}"
            )
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def _has_schema_error(errors: Iterable[dict[str, Any]]) -> bool:
    return any(error.get("type") != "value_error" or error.get("loc") for error in errors)


def _has_business_rule_error(errors: Iterable[dict[str, Any]]) -> bool:
    return any(error.get("type") == "value_error" and not error.get("loc") for error in errors)


def _schema_reason(errors: list[dict[str, Any]]) -> str:
    first = errors[0] if errors else {}
    location = _format_location(first.get("loc"))
    error_type = str(first.get("type", "schema_error"))
    return f"字段校验失败：{location or 'unknown'} / {error_type}"


def _business_rule_reason(errors: list[dict[str, Any]]) -> str:
    first = next((error for error in errors if error.get("type") == "value_error" and not error.get("loc")), errors[0] if errors else {})
    message = str(first.get("msg", "业务规则校验失败"))
    if message.lower().startswith("value error, "):
        message = message.split(", ", 1)[1]
    return f"业务规则校验失败：{message}"


def _format_location(location: Any) -> str:
    if location is None:
        return ""
    if isinstance(location, tuple):
        return ".".join(str(part) for part in location) if location else ""
    if isinstance(location, list):
        return ".".join(str(part) for part in location)
    return str(location)


def _shorten(text: str, limit: int = 120) -> str:
    cleaned = " ".join(str(text).split())
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[: limit - 1].rstrip() + "…"


def _selected_code_matches_expected(row: TrialResultRow) -> bool:
    expected_code = (row.expected_company_code or "").strip()
    if not expected_code:
        return False
    return (row.selected_company_code or "").strip() == expected_code


def _selected_code_mismatches_expected(row: TrialResultRow) -> bool:
    expected_code = (row.expected_company_code or "").strip()
    if not expected_code:
        return False
    return (row.selected_company_code or "").strip() != expected_code


def _is_unsafe_auto_code(row: TrialResultRow) -> bool:
    if not _row_can_auto_code(row):
        return False
    if row.expected_result_status != "strong_auto_code":
        return True
    if _selected_code_mismatches_expected(row):
        return True
    return False


def _to_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, int | float):
        return value != 0
    if isinstance(value, str):
        return value.strip().lower() in {"true", "1", "yes", "y"}
    return False


def _row_can_auto_code(row: TrialResultRow) -> bool:
    if _to_bool(row.can_auto_code):
        return True
    raw_output = (row.raw_model_output or "").strip()
    if not raw_output:
        return False
    try:
        payload = json.loads(raw_output)
    except json.JSONDecodeError:
        return False
    return _to_bool(payload.get("can_auto_code"))
